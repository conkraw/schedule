from helpers import login_for_test
"""Exact-date reporting regression tests. Invented provider/student data only."""
import copy
import csv
from datetime import date, timedelta
from io import BytesIO, StringIO
import json
from pathlib import Path
import unittest
from zipfile import ZipFile

from helpers import st, FakeGitHub, Upload, secret_settings, run_app
from openpyxl import Workbook
from docx import Document
from schedule_app.services.opd_archive import GitHubOPDArchive, get_opd_archive_config, OPDArchiveError
from schedule_app.services.educational_time import TIME_CSV_COLUMNS, TIME_WORK_TYPE_CSV_COLUMNS
from schedule_app.services.reporting_periods import (
    ReportingPeriod, reporting_period_json, read_reporting_period_json,
    teaching_report_date_text,
)
from schedule_app.services.teaching_analysis import (
    teaching_scan_archives, teaching_filter_date_range, teaching_require_date_range_data,
    teaching_annual_rows, teaching_work_type_rows,
)
from schedule_app.services.learner_reach import LEARNER_REACH_COLUMNS
from schedule_app.reports.chair_summary import teaching_chair_summary_data
from schedule_app.reports.teaching_export import teaching_build_zip
from schedule_app.sections.reporting_date_controls import load_selected_preset
from schedule_app.services.reporting_presets import GitHubReportingPresets
from unittest.mock import patch
from schedule_app.settings import TEACHING_CSV_COLUMNS, TEACHING_WORK_TYPE_CSV_COLUMNS, TEACHING_CHAIR_SUMMARY_FILENAME


def small_opd(start, cells):
    wb = Workbook()
    wb.remove(wb.active)
    for site in dict.fromkeys(key[0] for key in cells):
        ws = wb.create_sheet(site)
        ws['A1'] = 'Site:'
        ws['B1'] = site
        for i, day in enumerate(('Monday','Tuesday','Wednesday','Thursday','Friday','Saturday','Sunday')):
            ws.cell(3, i+2, day)
            ws.cell(4, i+2, start+timedelta(days=i)).number_format='mm/dd/yyyy'
        ws['A6'] = 'AM'
        ws['A7'] = 'AM'
        ws['A8'] = 'PM'
    for (site, address), value in cells.items():
        wb[site][address] = value
    output=BytesIO()
    wb.save(output)
    return output.getvalue()


def fixture_archive():
    secrets=secret_settings()
    st.reset(secrets=secrets)
    repo=FakeGitHub()
    client=GitHubOPDArchive(get_opd_archive_config(),transport=repo)
    client.save(small_opd(date(2026,2,16), {
        ('HOPE_DRIVE','B6'): 'Adams, Alex ~ Before Start',
        ('HOPE_DRIVE','C6'): 'Adams, Alex ~ Learner One',
        ('HOPE_DRIVE','C7'): 'Adams, Alex ~ Learner Two',
        ('ETOWN','C6'): 'Adams, Alex ~ Learner One',
        ('NYES','D6'): 'Brown, Blair ~ Learner Three',
        ('HOPE_DRIVE','E6'): 'Available, Only ~ ',
    }))
    client.save(small_opd(date(2026,6,29), {
        ('WARD A','C6'): 'Adams, Alex ~ Learner One',
        ('COMPLEX','D6'): 'Adams, Alex ~ Learner One',
        ('WARD A','E6'): 'Brown, Blair ~ Learner Two',
        ('WARD A','E7'): 'Brown, Blair ~ Learner Two', # benign same-work-type duplicate
    }))
    client.save(small_opd(date(2027,3,15), {
        ('NYES','C6'): 'Adams, Alex ~ Learner One',
        ('NYES','C7'): 'Adams, Alex ~ Learner Two',
        ('HOPE_DRIVE','D6'): 'Adams, Alex ~ After End',
        ('PSHCH_NURSERY','B6'): 'Diaz, Drew ~ Learner Three',
    }))
    client.save(small_opd(date(2027,3,22), {
        ('SJR_HOSP','B6'): 'SJR_1 ~ Outside Learner',
        ('SJR_HOSP','C6'): ' ~ Missing Provider',
    }))
    return secrets,repo,client


def all_doc_text(raw):
    doc=Document(BytesIO(raw))
    return '\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])


class ReportingDatesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.secrets,cls.repo,cls.client=fixture_archive()
        cls.scan=teaching_scan_archives(cls.client)
        cls.period=ReportingPeriod('26-27',date(2026,2,17),date(2027,3,16))
        cls.view=teaching_filter_date_range(cls.scan,cls.period)
        cls.zip_bytes,cls.rows=teaching_build_zip(cls.view,[2026])

    def test_daily_scan_keeps_no_student_identifiers(self):
        encoded=json.dumps(self.scan)
        for word in ('Learner One','Learner Two','Before Start','After End','Outside Learner','Missing Provider'):
            self.assertNotIn(word,encoded)
        teaching_require_date_range_data(self.scan)

    def test_exact_inclusive_boundaries_and_partial_rotations(self):
        self.assertEqual(sum(r['no_of_shifts'] for r in self.rows),9)
        self.assertEqual({r['preceptor_name']:r['no_of_shifts'] for r in self.rows},
                         {'Adams, Alex':6,'Brown, Blair':2,'Diaz, Drew':1})
        self.assertEqual(sum(r['educational_hours'] for r in self.rows),28)

    def test_july_does_not_split_custom_period(self):
        self.assertEqual({r['academic_year'] for r in self.rows},{'26-27'})
        self.assertEqual({r['academic_start_year'] for r in self.view['monthly']},{2026})
        self.assertEqual(len(self.rows),3) # only teaching contributors are listed
        self.assertEqual(len(teaching_chair_summary_data(self.view,[2026])),1)

    def test_same_day_counts_both_students(self):
        view=teaching_filter_date_range(self.scan,ReportingPeriod('One day',date(2026,2,17),date(2026,2,17)))
        rows=teaching_annual_rows(view,[2026])
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['no_of_shifts'],2)
        self.assertEqual(rows[0]['educational_hours'],4)

    def test_custom_label_is_not_derived_from_start_year(self):
        view=teaching_filter_date_range(self.scan,ReportingPeriod('Teaching cohort 2027',self.period.start,self.period.end))
        self.assertEqual({r['academic_year'] for r in teaching_annual_rows(view,[2026])},{'Teaching cohort 2027'})

    def test_boundary_exclusion_one_day_later(self):
        view=teaching_filter_date_range(self.scan,ReportingPeriod('26-27',date(2026,2,18),date(2027,3,15)))
        self.assertEqual(sum(r['no_of_shifts'] for r in teaching_annual_rows(view,[2026])),5)

    def test_more_than_twelve_months_allowed(self):
        self.assertGreater((self.period.end-self.period.start).days,365)
        self.assertEqual(len(teaching_annual_rows(self.view,[2026])),3)

    def test_mid_month_filtered_before_month_totals(self):
        alex=[r for r in self.view['monthly'] if r['preceptor_name']=='Adams, Alex']
        self.assertEqual({r['month']:r['no_of_shifts'] for r in alex},
                         {'2026-02-01':2,'2026-06-01':1,'2026-07-01':1,'2027-03-01':2})

    def test_work_type_subtotals_reconcile(self):
        typed=teaching_work_type_rows(self.view,[2026])
        self.assertEqual(sum(r['no_of_shifts'] for r in typed),9)
        self.assertEqual(sum(r['educational_hours'] for r in typed),28)
        self.assertEqual({r['work_type'] for r in typed},
                         {'Academic Pediatrics','Ward A','PSHCH Nursery','Complex Care'})

    def test_cross_type_duplicates_block_within_custom_dates(self):
        from schedule_app.services.teaching_validation import TeachingConflictError
        from test_learner_reach import scan_cells
        scan,*_=scan_cells({('NYES','B6'):'Example, Provider ~ Student One',
                           ('COMPLEX','B6'):'Example, Provider ~ Student One'})
        view=teaching_filter_date_range(scan,ReportingPeriod('Custom cohort',date(2026,9,7),date(2026,9,8)))
        self.assertEqual(view['work_type_conflicts'][0]['academic_year'],'Custom cohort')
        with self.assertRaises(TeachingConflictError):
            teaching_build_zip(view,[2026])
        self.assertEqual(self.scan['duplicate_assignments_removed'],2)


    def test_aggregate_scan_not_changed_by_filters(self):
        before=copy.deepcopy(self.scan)
        view=teaching_filter_date_range(self.scan,self.period)
        view['daily_by_work_type'][0]['source_sites'].append('Testing')
        self.assertEqual(self.scan,before)
        self.assertNotIn('reporting_period',self.scan)

    def test_cannot_re_filter_subset_as_full_scan(self):
        with self.assertRaises(OPDArchiveError):
            teaching_filter_date_range(self.view,self.period)

    def test_blank_and_invalid_inputs_rejected(self):
        bad=[('',self.period.start,self.period.end),('x',None,self.period.end),('x',self.period.end,self.period.start),
             ('x'*61,self.period.start,self.period.end),('bad\x00',self.period.start,self.period.end),
             ('bad\uffff',self.period.start,self.period.end),('x',date(1969,1,1),self.period.end)]
        for values in bad:
            with self.subTest(values=values), self.assertRaises(OPDArchiveError):ReportingPeriod(*values)

    def test_leap_day_and_same_day_valid(self):
        period=ReportingPeriod('Leap-day',date(2028,2,29),date(2028,2,29))
        self.assertEqual(period.start,period.end)

    def test_empty_period_does_not_make_empty_report(self):
        view=teaching_filter_date_range(self.scan,ReportingPeriod('Empty',date(2028,1,1),date(2028,2,1)))
        self.assertEqual(teaching_annual_rows(view,[2028]),[])
        with self.assertRaises(OPDArchiveError):teaching_build_zip(view,[2028])

    def test_old_monthly_only_scan_rejected(self):
        old=dict(self.scan)
        old.pop('daily_by_work_type')
        with self.assertRaisesRegex(OPDArchiveError,'refresh'):
            teaching_filter_date_range(old,self.period)

    def test_mismatched_daily_totals_rejected(self):
        bad=copy.deepcopy(self.scan)
        bad['daily_by_work_type'][0]['no_of_shifts']+=1
        bad['daily_by_work_type'][0]['educational_hours']+=4
        with self.assertRaisesRegex(OPDArchiveError,'Daily teaching counts'):
            teaching_filter_date_range(bad,self.period)

    def test_json_roundtrip_and_validation(self):
        self.assertEqual(read_reporting_period_json(reporting_period_json(self.period)),self.period)
        for raw in (b'not json',b'[]',b'{}',b'{"reporting_period_version":99}',b'X'*8200):
            with self.subTest(raw=raw[:25]),self.assertRaises(OPDArchiveError):read_reporting_period_json(raw)

    def test_filename_label_cannot_create_paths(self):
        period=ReportingPeriod('../../chair/special',self.period.start,self.period.end)
        self.assertNotIn('/',period.filename_part())
        self.assertIn('2026-02-17_to_2027-03-16',period.filename_part())

    def test_csv_columns_and_period_metadata(self):
        with ZipFile(BytesIO(self.zip_bytes)) as z:
            self.assertIsNone(z.testzip())
            for file,headers in [('preceptor_teaching_summary.csv',TIME_CSV_COLUMNS),
                                 ('preceptor_teaching_by_work_type.csv',TIME_WORK_TYPE_CSV_COLUMNS)]:
                reader=csv.DictReader(StringIO(z.read(file).decode('utf-8-sig')))
                self.assertEqual(tuple(reader.fieldnames),tuple(headers))
                self.assertEqual({r['academic_year'] for r in reader},{'26-27'})
            self.assertEqual(read_reporting_period_json(z.read('Reporting_Period.json')),self.period)
            meta=list(csv.DictReader(StringIO(z.read('Reporting_Period.csv').decode('utf-8-sig'))))[0]
            self.assertEqual(meta['start_date'],'2026-02-17')
            self.assertEqual(meta['end_date'],'2027-03-16')
            self.assertEqual(meta['both_dates_included'],'YES')

    def test_word_reports_show_true_dates_not_july_dates(self):
        with ZipFile(BytesIO(self.zip_bytes)) as z:
            for filename in [f for f in z.namelist() if f.endswith('.docx')]:
                text=all_doc_text(z.read(filename))
                self.assertIn('Reporting period 26-27',text)
                self.assertIn('February 17, 2026 - March 16, 2027',text)
                self.assertNotIn('July 1, 2026 - June 30, 2027',text)
                self.assertEqual(text.count('Reporting period 26-27'),1)
                for student in ('Learner One','Learner Two','Outside Learner'):self.assertNotIn(student,text)

    def test_archive_source_coverage_and_unresolved_outside_range(self):
        data=teaching_chair_summary_data(self.view,[2026])[0]
        self.assertEqual(data['source_count'],3)
        self.assertEqual(data['named_preceptor_count'],3)
        self.assertEqual(data['named_preceptors_with_students'],3)
        self.assertEqual(data['unresolved_labels'],[])
        self.assertFalse(data['has_missing_provider'])
        with ZipFile(BytesIO(self.zip_bytes)) as z:
            notes=z.read('Report_Notes.txt').decode()
            self.assertIn('Exact reporting dates: 2026-02-17 through 2027-03-16',notes)
            self.assertNotIn('Academic year is July 1',notes)
            self.assertIn('archive-wide',notes)

    def test_legacy_multi_year_mode_preserved(self):
        blob,rows=teaching_build_zip(self.scan,[2025,2026])
        self.assertEqual({r['academic_year'] for r in rows},{'25-26','26-27'})
        with ZipFile(BytesIO(blob)) as z:
            self.assertNotIn('Reporting_Period.json',z.namelist())
            chair=all_doc_text(z.read(TEACHING_CHAIR_SUMMARY_FILENAME))
            self.assertIn('July 1, 2025 - June 30, 2026',chair)
            self.assertIn('July 1, 2026 - June 30, 2027',chair)

    def run_ui(self,extra=None,state=None):
        values={'schedule_app_mode':'PTS',
                'teaching_period_start':self.period.start,'teaching_period_end':self.period.end,
                'teaching_period_label':self.period.label}
        values.update(extra or {})
        return run_app(values,secrets=self.secrets,state=state,repo=self.repo, evaluation_login=True)

    def test_ui_builds_custom_zip(self):
        result=self.run_ui({'teaching_load_archives':True,'teaching_build_zip':True})
        self.assertIn('Preceptor_Teaching_26-27_2026-02-17_to_2027-03-16.zip',result['downloads'])
        self.assertIn(TEACHING_CHAIR_SUMMARY_FILENAME,result['downloads'])

    def test_ui_date_edit_discards_old_zip_and_does_not_reload(self):
        first=self.run_ui({'teaching_load_archives':True,'teaching_build_zip':True})
        second=self.run_ui({'teaching_period_start':date(2026,2,18)},first['state'])
        self.assertNotIn('teaching_zip',second['state'])
        self.assertEqual(first['state']['teaching_scan'],second['state']['teaching_scan'])
        self.assertFalse(any(name.endswith('.zip') for name in second['downloads']))

    def test_ui_label_edit_discards_old_zip(self):
        first=self.run_ui({'teaching_load_archives':True,'teaching_build_zip':True})
        second=self.run_ui({'teaching_period_label':'Chair review'},first['state'])
        self.assertNotIn('teaching_zip',second['state'])

    def test_ui_invalid_dates_hide_old_downloads(self):
        first=self.run_ui({'teaching_load_archives':True,'teaching_build_zip':True})
        result=self.run_ui({'teaching_period_end':date(2020,1,1)},first['state'])
        self.assertNotIn('teaching_zip',result['state'])
        self.assertFalse(any(name.endswith('.zip') for name in result['downloads']))

    def test_ui_preferences_survive_widget_cleanup(self):
        result=self.run_ui()
        state=result['state']
        for key in ('teaching_period_start','teaching_period_end','teaching_period_label','teaching_reporting_mode'):
            state.pop(key,None)
        again=run_app({'schedule_app_mode':'PTS'},secrets=self.secrets,state=state,repo=self.repo, evaluation_login=True)
        self.assertEqual(again['state']['teaching_period_start'],self.period.start)
        self.assertEqual(again['state']['teaching_period_label'],self.period.label)

    def test_ui_load_github_period_callback(self):
        # The local JSON uploader was deliberately replaced with a GitHub preset.
        service=GitHubReportingPresets(self.client)
        initial=service.load()
        receipt=service.save('Callback example',self.period,initial)
        st.reset(secrets=self.secrets,state={'teaching_presets_snapshot':receipt['snapshot'],
                                           'teaching_zip':b'old'})
        login_for_test()
        with patch('schedule_app.services.opd_archive.requests.request',self.repo.request):
            load_selected_preset(receipt['preset_id'])
        self.assertEqual(st.session_state['teaching_period_start'],self.period.start)
        self.assertEqual(st.session_state['teaching_period_label'],'26-27')
        self.assertNotIn('teaching_zip',st.session_state)

    def test_ui_defaults_do_not_guess_user_dates(self):
        result=run_app({'schedule_app_mode':'PTS'},secrets=self.secrets,repo=self.repo, evaluation_login=True)
        self.assertIsNone(result['state']['teaching_period_start'])
        self.assertIsNone(result['state']['teaching_period_end'])
        self.assertEqual(result['state']['teaching_reporting_mode'],'Custom dates')
        self.assertNotIn('teaching_zip',result['state'])

if __name__=='__main__':
    unittest.main()

# Time-based CSV schema supersedes the old student-weighted export contract.
