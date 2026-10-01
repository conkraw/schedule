"""Selectable assessment shifts: encryption, cross-session defaults and reports.

All people/data/repositories are invented; requests use in-memory test doubles.
"""
from copy import deepcopy
from dataclasses import replace
from datetime import date
from io import BytesIO, StringIO
import csv
import json
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from cryptography.fernet import Fernet
from docx import Document
from helpers import FakeGitHub, FakeResponse, st, secret_settings, login_for_test
from test_learner_reach import scan_cells
from test_assessment_completion import NAME, record, raw_csv
from schedule_app.services.opd_archive import OPDArchiveConfig, GitHubOPDArchive, OPDArchiveError
from schedule_app.services.assessment_settings import (
    DEFAULT_MINIMUM_SHIFTS, MAX_MINIMUM_SHIFTS, FILENAME,
    GitHubAssessmentSettings, validate_minimum_shifts,
)
from schedule_app.sections import assessment_settings as ui
from schedule_app.sections import assessment_completion as page
from schedule_app.services.assessment_completion import (
    ASSESSMENT_VERSION, HP, COLUMNS, load_completion_inputs, build_completion_bundle,
    unverified_bundle, completion_signature, completion_rows, assessment_method_note,
)
from schedule_app.services.oasis_student_evaluations import GitHubOASISStudentEvaluations
from schedule_app.services.preceptor_oasis_links import GitHubPreceptorOASISLinks
from schedule_app.services.student_name_review import student_name_review, review_table_rows
from schedule_app.services.reporting_periods import ReportingPeriod
from schedule_app.services.teaching_analysis import teaching_filter_date_range, teaching_annual_rows
from schedule_app.services.student_continuity import student_continuity_counts
from schedule_app.services.evaluation_access import lock_evaluation_records
from schedule_app.reports.assessment_completion import append_individual_completion, append_chair_completion
from schedule_app.reports.teaching_export import teaching_build_zip


@pytest.fixture
def storage():
    secrets = secret_settings()
    st.reset(secrets=secrets)
    login_for_test()
    repo = FakeGitHub()
    archive = GitHubOPDArchive(OPDArchiveConfig(**secrets['opd_archive']), transport=repo)
    return GitHubAssessmentSettings(archive), repo, secrets


def show(storage, values=None, state=None, *, authenticated=True):
    service, repo, secrets = storage
    st.reset(values=values, state=state, secrets=secrets)
    if authenticated:
        login_for_test()
    value = ui.render_assessment_threshold(service.archive)
    return value, dict(st.session_state)


def test_initial_default_three_without_writing(storage):
    store, repo, _ = storage
    snap = store.load()
    assert snap['minimum_shifts'] == 3 and snap['sha'] is None
    assert repo.write_count == 0
    assert show(storage)[0] == 3
    assert repo.write_count == 0


@pytest.mark.parametrize('value', [1, 3, 4, 5, 10, MAX_MINIMUM_SHIFTS])
def test_saved_value_reloads_in_new_service(storage, value):
    store, repo, _ = storage
    saved = store.save(value, expected=store.load())
    another = GitHubAssessmentSettings(store.archive).load()
    assert another == saved
    assert another['minimum_shifts'] == value
    token = repo.tree[store.path]
    assert b'minimum_shifts' not in token
    data = json.loads(store.cipher.decrypt(token))
    assert data == {'kind': 'pts_assessment_settings', 'version': 1, 'minimum_shifts': value}
    assert store.path == 'opd_archive/' + FILENAME


@pytest.mark.parametrize('value', [0, -1, True, False, 3.0, 3.5, '4', None, MAX_MINIMUM_SHIFTS + 1])
def test_invalid_values_do_not_write(storage, value):
    store, repo, _ = storage
    with pytest.raises(OPDArchiveError):
        store.save(value, expected=store.load())
    assert repo.write_count == 0


def test_same_saved_value_does_not_write_again(storage):
    store, repo, _ = storage
    saved = store.save(4, expected=store.load())
    store.save(4, expected=saved)
    assert repo.write_count == 1


def test_stale_session_cannot_overwrite_newer_setting(storage):
    store, repo, _ = storage
    first = store.load()
    store.save(4, expected=first)
    with pytest.raises(OPDArchiveError, match='another session'):
        store.save(5, expected=first)
    assert store.load()['minimum_shifts'] == 4
    assert repo.write_count == 1


def test_other_repository_scope_cannot_write(storage):
    store, repo, _ = storage
    first = store.load(); first['scope'] = 'another archive'
    with pytest.raises(OPDArchiveError, match='Load'):
        store.save(5, expected=first)
    assert repo.write_count == 0


def test_unrelated_opd_commit_and_presets_unchanged(storage):
    store, repo, _ = storage
    first = store.load()
    others = {'opd_archive/OPD_2026-09-07.xlsx.enc': b'unchanged OPD',
              'opd_archive/reporting_date_presets.json.enc': b'unchanged dates',
              'opd_archive/student_assessment_id_links.json.enc': b'unchanged student links'}
    repo.tree.update(others); repo._commit()
    store.save(5, expected=first)
    assert all(repo.tree[k] == v for k, v in others.items())
    assert store.archive.list_rotations() == [date(2026, 9, 7)]


def test_wrong_key_fails_without_reset_or_overwrite(storage):
    store, repo, _ = storage
    store.save(5, expected=store.load())
    bad = replace(store.config, encryption_key=Fernet.generate_key().decode())
    with pytest.raises(OPDArchiveError, match='verified/decrypted'):
        GitHubAssessmentSettings(GitHubOPDArchive(bad, transport=repo)).load()
    assert repo.write_count == 1


@pytest.mark.parametrize('data', [b'{"kind":"pts_assessment_settings","version":1,"minimum_shifts":4,"minimum_shifts":5}',
                                   b'bad-json', b'{"kind":"pts_assessment_settings","version":1,"minimum_shifts":0}',
                                   b'{"kind":"pts_assessment_settings","version":2,"minimum_shifts":5}'])
def test_corrupt_or_unknown_settings_do_not_default(storage, data):
    store, repo, _ = storage
    repo.tree[store.path] = store.cipher.encrypt(data); repo._commit()
    with pytest.raises(OPDArchiveError):
        store.load()
    assert repo.write_count == 0


def test_write_response_without_commit_not_claimed_saved(storage):
    store, repo, _ = storage
    actual = repo.request
    def broken(method, url, **kwargs):
        if method == 'PUT':
            return FakeResponse(201, {})
        return actual(method, url, **kwargs)
    with patch.object(repo, 'request', side_effect=broken):
        with pytest.raises(OPDArchiveError, match='not confirmed'):
            store.save(5, expected=store.load())


def test_wrong_readback_is_not_success(storage):
    store, repo, _ = storage
    actual = store.load
    def wrong(*, commit=None):
        result = actual(commit=commit)
        if commit:
            result['minimum_shifts'] = 6
        return result
    with patch.object(store, 'load', side_effect=wrong):
        with pytest.raises(OPDArchiveError, match='could not be verified'):
            store.save(5, expected=store.load())


def test_number_change_autosaves_and_next_session_reuses_it(storage):
    _, state = show(storage)
    value, state = show(storage, {ui.P + 'value': 4}, state)
    assert value == 4 and storage[0].load()['minimum_shifts'] == 4
    assert show(storage)[0] == 4
    value, state = show(storage, {ui.P + 'value': 5}, state)
    assert value == 5 and show(storage)[0] == 5
    assert storage[1].write_count == 2


def test_regular_widget_rerun_does_not_write(storage):
    _, state = show(storage)
    _, state = show(storage, {ui.P + 'value': 4}, state)
    n = storage[1].write_count
    assert show(storage, state=state)[0] == 4
    assert storage[1].write_count == n


def test_switching_sections_restores_deleted_widget_key(storage):
    _, state = show(storage)
    _, state = show(storage, {ui.P + 'value': 5}, state)
    del state[ui.P + 'value']  # Streamlit drops keys of widgets not drawn.
    assert show(storage, state=state)[0] == 5


def test_saving_clears_reports_but_keeps_loaded_records_and_dates(storage):
    _, state = show(storage)
    state.update(teaching_zip=b'old-report', teaching_zip_signature='old',
                 teaching_scan={'same OPD': True}, assessment_completion_inputs={'keep': True},
                 teaching_period_label='26-27')
    value, next_state = show(storage, {ui.P + 'value': 4}, state)
    assert value == 4
    assert 'teaching_zip' not in next_state and 'teaching_zip_signature' not in next_state
    assert next_state['teaching_scan'] == state['teaching_scan']
    assert next_state['assessment_completion_inputs'] == state['assessment_completion_inputs']
    assert next_state['teaching_period_label'] == '26-27'


def test_load_failure_not_reset_three_and_explicit_retry(storage):
    store, repo, _ = storage
    store.save(5, expected=store.load())
    with patch.object(GitHubAssessmentSettings, 'load', side_effect=OPDArchiveError('offline')) as load:
        value, state = show(storage)
        assert value is None
        value, state = show(storage, state=state)
        assert value is None and load.call_count == 1
    assert not any(k == ui.P + 'value' for _, k in st.widget_keys)
    assert show(storage, {ui.P + 'reload': True}, state)[0] == 5


def test_save_failure_not_retried_on_unrelated_rerun(storage):
    _, state = show(storage)
    with patch.object(GitHubAssessmentSettings, 'save', side_effect=OPDArchiveError('offline')) as save:
        value, state = show(storage, {ui.P + 'value': 4}, state)
        assert value is None
        value, state = show(storage, state=state)
        assert value is None and save.call_count == 1
    value, state = show(storage, {ui.P + 'retry': True}, state)
    assert value == 4 and storage[0].load()['minimum_shifts'] == 4


def test_simultaneous_changes_require_review_then_reload(storage):
    store, repo, _ = storage
    _, state = show(storage)
    store.save(5, expected=store.load())
    value, state = show(storage, {ui.P + 'value': 4}, state)
    assert value is None and store.load()['minimum_shifts'] == 5
    assert show(storage, {ui.P + 'reload': True}, state)[0] == 5


def test_locked_section_cannot_read_or_save(storage):
    store, repo, _ = storage
    with patch.object(GitHubAssessmentSettings, 'load') as load, patch.object(GitHubAssessmentSettings, 'save') as save:
        assert show(storage, {ui.P + 'value': 5}, authenticated=False)[0] is None
        load.assert_not_called(); save.assert_not_called()
    assert not st.widget_keys


def test_lock_clears_local_value_but_not_github(storage):
    _, state = show(storage)
    _, state = show(storage, {ui.P + 'value': 5}, state)
    lock_evaluation_records()
    assert not any(str(k).startswith(ui.P) for k in st.session_state)
    assert storage[0].load()['minimum_shifts'] == 5
    assert show(storage)[0] == 5


@pytest.fixture
def dataset():
    # Alpha 3 shifts; Beta 4; Gamma 5. AM/PM on a single date remain distinct.
    locations = ['B6','B8','C6','D6','E6']
    counts = {'Alpha':3, 'Beta':4, 'Gamma':5}
    cells = {('NYES', pos): NAME + ' ~ ' + '; '.join('Learner, ' + who for who, n in counts.items() if n > i)
             for i, pos in enumerate(locations)}
    cells['NYES','F6'] = NAME + ' ~ '
    scan, repo, archive, secrets, raw = scan_cells(cells)
    view = teaching_filter_date_range(scan, ReportingPeriod('26-27', date(2026,9,7), date(2026,9,13)))
    links = GitHubPreceptorOASISLinks(archive)
    links.save_username(NAME, 'avery', expected=links.load())
    source = GitHubOASISStudentEvaluations(archive)
    source.save(raw_csv([record(form=HP),
                         record('2', sid='id-beta', student='Learner, Beta'),
                         record('3', sid='id-gamma', student='Learner, Gamma')]))
    inputs = load_completion_inputs(archive, view, [2026])
    inputs['summaries']['2026'] = {'rows_by_id': {'avery': {'evaluation_count': 2}}}
    return dict(scan=scan, view=view, repo=repo, archive=archive, secrets=secrets, inputs=inputs)


@pytest.mark.parametrize('minimum,eligible,clinical,hp', [(3,3,66.7,33.3),(4,2,100.0,0.0),(5,1,100.0,0.0),(6,0,None,None)])
def test_threshold_changes_both_numerator_membership_and_denominator(dataset, minimum, eligible, clinical, hp):
    bundle, missing = build_completion_bundle(dataset['inputs'], dataset['view'], [2026], minimum_shifts=minimum)
    row = bundle['rows'][0]
    assert row['eligible_students'] == eligible
    assert row['clinical_completion_pct'] == clinical
    assert row['hp_completion_pct'] == hp
    assert row['minimum_shifts'] == bundle['minimum_shifts'] == minimum
    assert not missing
    assert row['clinical_forms_submitted'] == 2  # All submitted forms audit count unchanged.
    if eligible == 0:
        assert row['assessment_status'] == 'No eligible students (6+ shifts)'


def test_default_still_three_and_input_not_mutated(dataset):
    before = deepcopy(dataset['inputs'])
    result, _ = build_completion_bundle(dataset['inputs'], dataset['view'], [2026])
    assert result['minimum_shifts'] == 3 and result['rows'][0]['eligible_students'] == 3
    assert dataset['inputs'] == before


def test_calendar_days_measure_unaffected(dataset):
    before = student_continuity_counts(dataset['view'], NAME, 2026)
    hours = teaching_annual_rows(dataset['view'], [2026])
    for minimum in [3,4,5]:
        build_completion_bundle(dataset['inputs'], dataset['view'], [2026], minimum_shifts=minimum)
    assert before == student_continuity_counts(dataset['view'], NAME, 2026)
    assert before['unique_students_3plus_days'] == 2
    assert teaching_annual_rows(dataset['view'], [2026]) == hours


def test_duplicates_do_not_help_reach_four_shifts(dataset):
    inputs = deepcopy(dataset['inputs'])
    inputs['assignments'] += [r for r in inputs['assignments'] if 'Alpha' in r['student']] * 3
    bundle, _ = build_completion_bundle(inputs, dataset['view'], [2026], minimum_shifts=4)
    assert bundle['rows'][0]['eligible_students'] == 2


def test_days_outside_range_do_not_help_reach_four_shifts(dataset):
    inputs = deepcopy(dataset['inputs'])
    inputs['assignments'].append({**inputs['assignments'][0], 'student':'Learner, Alpha', 'date':'2026-09-14'})
    bundle, _ = build_completion_bundle(inputs, dataset['view'], [2026], minimum_shifts=4)
    assert bundle['rows'][0]['eligible_students'] == 2


def test_unmatched_flags_follow_threshold_not_fixed_three(dataset):
    inputs = deepcopy(dataset['inputs'])
    inputs['prepared']['name_ids'].pop('learner, alpha')
    bundle, unmatched = build_completion_bundle(inputs, dataset['view'], [2026], minimum_shifts=3)
    assert len(unmatched) == 1 and bundle['rows'][0]['clinical_completion_pct'] == 66.7
    assert bundle['rows'][0]['assessment_status'].startswith('Provisional:')
    bundle, unmatched = build_completion_bundle(inputs, dataset['view'], [2026], minimum_shifts=4)
    assert not unmatched and bundle['rows'][0]['clinical_completion_pct'] == 100
    review = student_name_review(inputs, dataset['view'], [2026], unmatched)
    assert review['eligible_missing_count'] == 0 and 'learner, alpha' in review['missing']
    rows = review_table_rows(review['missing'], minimum_shifts=4)
    assert rows[0]['Affects completion percentage'] == 'NO — below 4-shift threshold'


def test_report_signature_changes_even_if_result_has_same_percentage(dataset):
    a,_ = build_completion_bundle(dataset['inputs'], dataset['view'], [2026], minimum_shifts=4)
    b,_ = build_completion_bundle(dataset['inputs'], dataset['view'], [2026], minimum_shifts=5)
    assert completion_signature(a) != completion_signature(b)


def test_report_rejects_mixed_thresholds(dataset):
    b,_ = build_completion_bundle(dataset['inputs'], dataset['view'], [2026], minimum_shifts=4)
    b['rows'][0]['minimum_shifts'] = 3
    with pytest.raises(OPDArchiveError, match='inconsistent'):
        completion_rows(b, dataset['view'], 2026)


def test_cached_previous_version_rejected(dataset):
    b = unverified_bundle(dataset['view'], [2026]); b['version'] = ASSESSMENT_VERSION - 1
    with pytest.raises(OPDArchiveError):
        completion_rows(b, dataset['view'], 2026)


def text(doc):
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


@pytest.mark.parametrize('minimum',[3,4,5])
def test_chair_individual_csv_notes_all_show_selected_threshold(dataset, minimum):
    b,_ = build_completion_bundle(dataset['inputs'], dataset['view'], [2026], minimum_shifts=minimum)
    view = dict(dataset['view'], assessment_completion=b)
    individual = Document(); append_individual_completion(individual, view, NAME, 2026)
    chair = Document(); append_chair_completion(chair, view, 2026)
    assert f'Students assigned for {minimum}+ shifts' in text(individual)
    assert f'Students\n{minimum}+ shifts' in text(chair)
    for doc in (individual,chair):
        assert f'at least {minimum} distinct AM/PM shifts' in text(doc)
        assert 'Learner, Alpha' not in text(doc) and 'id-alpha' not in text(doc)
    raw,_ = teaching_build_zip(view,[2026])
    with ZipFile(BytesIO(raw)) as z:
        rows = list(csv.DictReader(StringIO(z.read('preceptor_student_assessment_completion.csv').decode('utf-8-sig'))))
        assert rows[0]['minimum_shifts'] == str(minimum)
        assert 'eligible_students' in rows[0] and 'eligible_students_3plus_shifts' not in rows[0]
        assert f'at least {minimum} distinct AM/PM shifts' in z.read('Report_Notes.txt').decode()


def test_unknown_setting_yields_unchecked_not_three(dataset):
    b = unverified_bundle(dataset['view'], [2026], 'Not checked: settings unavailable', minimum_shifts=None)
    assert b['minimum_shifts'] is None and b['rows'][0]['minimum_shifts'] is None
    doc = Document(); append_chair_completion(doc, dict(dataset['view'], assessment_completion=b),2026)
    assert 'minimum-shifts setting was not verified' in text(doc)
    assert 'Students\n3+ shifts' not in text(doc)


def test_loaded_records_reused_when_number_changes(dataset):
    secrets = dataset['secrets']
    st.reset(values={page.P + 'refresh':True}, secrets=secrets)
    login_for_test()
    with patch.object(page, 'load_completion_inputs', return_value=dataset['inputs']) as load:
        b = page.render_assessment_completion(dataset['archive'], dataset['view'],[2026])
        assert b['minimum_shifts'] == 3 and load.call_count == 1
        state = dict(st.session_state)
        state['teaching_zip'] = b'old'
        st.reset(values={ui.P + 'value':4},state=state,secrets=secrets); login_for_test()
        b = page.render_assessment_completion(dataset['archive'], dataset['view'],[2026])
        assert b['minimum_shifts'] == 4 and b['rows'][0]['eligible_students'] == 2
        assert load.call_count == 1 and 'teaching_zip' not in st.session_state
        assert any('Students 4+ shifts' in msg for _, msg in st.messages)


def test_unavailable_setting_prevents_calculation_but_not_teaching(dataset):
    st.reset(secrets=dataset['secrets']);login_for_test()
    with patch.object(GitHubAssessmentSettings, 'load', side_effect=OPDArchiveError('offline')):
        b = page.render_assessment_completion(dataset['archive'], dataset['view'],[2026])
    assert b['minimum_shifts'] is None and b['rows'][0]['clinical_completion_pct'] is None
    raw,_ = teaching_build_zip(dict(dataset['view'], assessment_completion=b),[2026])
    assert ZipFile(BytesIO(raw)).testzip() is None


def test_only_threshold_runtime_files_changed():
    from helpers import ROOT
    base = ROOT.parent / 'base'
    if (not (base / 'UPDATE_PTS_STUDENT_NAME_MATCHES.md').exists()
            or (base / 'UPDATE_PTS_MINIMUM_SHIFTS.md').exists()):
        pytest.skip('Release packaging check requires the supplied previous-version baseline')
    changed = {p.relative_to(ROOT).as_posix() for p in [ROOT/'app_sch_2026.py', *(ROOT/'schedule_app').rglob('*.py')]
               if not (base/p.relative_to(ROOT)).exists() or p.read_bytes() != (base/p.relative_to(ROOT)).read_bytes()}
    assert changed == {
        'schedule_app/services/assessment_settings.py',
        'schedule_app/sections/assessment_settings.py',
        'schedule_app/services/assessment_completion.py',
        'schedule_app/sections/assessment_completion.py',
        'schedule_app/services/student_name_review.py',
        'schedule_app/sections/student_name_matches.py',
        'schedule_app/reports/assessment_completion.py',
        'schedule_app/reports/teaching_export.py',
    }
