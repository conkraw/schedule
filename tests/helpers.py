"""Offline test doubles and invented OPD fixtures; no credentials or real names."""
import ast
import base64
from contextlib import contextmanager
from datetime import date, datetime, timedelta
import hashlib
import io
from pathlib import Path
import random
import runpy
import sys
from types import ModuleType
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook
from cryptography.fernet import Fernet

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

class StopRun(BaseException):
    pass

class FakeStreamlit(ModuleType):
    def __init__(self):
        super().__init__('streamlit')
        self.sidebar = self
        self.reset()

    def reset(self, *, values=None, secrets=None, state=None):
        self.values = dict(values or {})
        self.secrets = dict(secrets or {})
        self.session_state = dict(state or {})
        self.downloads = {}
        self.messages = []
        self.events = []
        self.widget_keys = []

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def _value(self, label, kwargs, default):
        key = kwargs.get('key') or label
        self.widget_keys.append((label, key))
        return self.values.get(key, default)

    def _message(self, name, *args, **kwargs):
        self.events.append(name)
        self.messages.append((name, str(args[0]) if args else ''))
        return self

    def set_page_config(self, **kwargs):
        self.events.append('set_page_config')

    def file_uploader(self, label, **kwargs):
        return self._value(label, kwargs, [] if kwargs.get('accept_multiple_files') else None)

    def text_input(self, label, value='', **kwargs):
        key = kwargs.get('key')
        selected = self._value(label, kwargs, self.session_state.get(key, value) if key else value)
        if key:
            self.session_state[key] = selected
        return selected

    def date_input(self, label, value=None, **kwargs):
        key = kwargs.get('key')
        selected = self._value(label, kwargs, self.session_state.get(key, value) if key else value)
        if key:
            self.session_state[key] = selected
        return selected

    def number_input(self, label, value="min", min_value=None, **kwargs):
        key = kwargs.get("key")
        initial = min_value if value == "min" and min_value is not None else (0 if value == "min" else value)
        selected = self._value(label, kwargs, self.session_state.get(key, initial) if key else initial)
        if key:
            self.session_state[key] = selected
        return selected

    def radio(self, label, options, index=0, **kwargs):
        key = kwargs.get('key')
        default = self.session_state.get(key, options[index]) if key else options[index]
        value = self._value(label, kwargs, default)
        if key:
            self.session_state[key] = value
        return value

    def selectbox(self, label, options, index=0, **kwargs):
        return self.radio(label, options, index=index, **kwargs)

    def multiselect(self, label, options, default=None, **kwargs):
        key = kwargs.get('key')
        default = default if default is not None else self.session_state.get(key, [])
        selected = self._value(label, kwargs, default)
        if key:
            self.session_state[key] = selected
        return selected

    def checkbox(self, label, value=False, **kwargs):
        return self._value(label, kwargs, value)

    toggle = checkbox

    def button(self, label, **kwargs):
        clicked = bool(self._value(label, kwargs, False)) and not kwargs.get('disabled', False)
        return clicked

    def download_button(self, label, data, file_name=None, mime=None, **kwargs):
        value = data.getvalue() if hasattr(data, 'getvalue') else data
        self.downloads[file_name or label] = value
        self.events.append('download_button')
        return False

    def columns(self, spec, **kwargs):
        return [self] * (spec if isinstance(spec, int) else len(spec))

    def cache_data(self, func=None, **kwargs):
        return func if func is not None else lambda f: f

    def stop(self):
        raise StopRun()

    def rerun(self):
        raise StopRun()

    def __getattr__(self, name):
        if name.startswith('__'):
            raise AttributeError(name)
        return lambda *args, **kwargs: self._message(name, *args, **kwargs)

# All modules use one mock object, reset for each test, so Python import caching
# also gets exercised. Real Streamlit and its UI server are not required.
st = FakeStreamlit()
sys.modules['streamlit'] = st

class Upload(io.BytesIO):
    def __init__(self, raw, name):
        super().__init__(raw)
        self.name = name
        self.size = len(raw)
        self.file_id = hashlib.sha256(raw).hexdigest()

class FakeResponse:
    def __init__(self, status=200, data=None, content=b''):
        self.status_code = status
        self.data = data
        self.content = content

    def json(self):
        return self.data

    def close(self):
        pass

    def iter_content(self, chunk_size):
        yield self.content

class FakeGitHub:
    def __init__(self):
        self.tree = {}
        self.snapshots = {}
        self.sequence = 0
        self.write_count = 0
        self._commit()

    @staticmethod
    def sha(raw):
        return hashlib.sha1(b'blob ' + str(len(raw)).encode() + b'\0' + raw).hexdigest()

    def _commit(self):
        self.sequence += 1
        self.head = f'{self.sequence:040x}'
        self.snapshots[self.head] = dict(self.tree)

    def request(self, method, url, *, params=None, json=None, **kwargs):
        if '/branches/' in url:
            return FakeResponse(data={'commit': {'sha': self.head}})
        path = unquote(url.split('/contents/', 1)[1])
        if method == 'DELETE':
            old = self.tree.get(path)
            if old is None:
                return FakeResponse(404)
            if json.get('sha') != self.sha(old):
                return FakeResponse(409)
            del self.tree[path]
            self._commit()
            self.write_count += 1
            return FakeResponse(200, {'commit': {'sha': self.head}})
        if method == 'PUT':
            old = self.tree.get(path)
            if old is not None and json.get('sha') != self.sha(old):
                return FakeResponse(409)
            self.tree[path] = base64.b64decode(json['content'])
            self._commit()
            self.write_count += 1
            return FakeResponse(201, {'commit': {'sha': self.head}})
        tree = self.snapshots[(params or {}).get('ref', self.head)]
        if path in tree:
            raw = tree[path]
            return FakeResponse(data={'type': 'file', 'name': path.rsplit('/', 1)[-1],
                                     'size': len(raw), 'sha': self.sha(raw),
                                     'encoding': 'base64', 'content': base64.b64encode(raw).decode()}, content=raw)
        entries = [{'type': 'file', 'name': name.rsplit('/',1)[-1]} for name in tree if name.startswith(path+'/')]
        return FakeResponse(data={'entries': entries}) if entries else FakeResponse(404)


def secret_settings(key=None):
    return {'opd_archive': {'owner': 'example-account', 'repo': 'opd-test-archive',
            'branch': 'main', 'folder': 'opd_archive', 'github_token': 'offline-test-token',
            'encryption_key': key or Fernet.generate_key().decode()}}

SITES = ['HOPE_DRIVE','ETOWN','NYES','LANCASTER','LANCASTER_CMG','COMPLEX','WARD A',
         'PSHCH_NURSERY','HAMPDEN_NURSERY','SJR_HOSP','AAC','AHOLOUKPE','ADOLMED']


def make_opd(start=date(2026,8,3), changes=None, availability_only=False):
    """Use the app's existing template, filling it with invented assignments."""
    from schedule_app.services.opd_workbooks import generate_opd_workbook
    dates = pd.DataFrame([{f'hd_day_date{i+1}': pd.Timestamp(start+timedelta(days=i)) for i in range(28)}])
    raw = generate_opd_workbook(dates)
    wb = load_workbook(io.BytesIO(raw))
    values = {
        ('HOPE_DRIVE','B8'): 'Adams, Alex ~ Learner One',
        ('HOPE_DRIVE','B9'): 'Adams, Alex ~ Learner Two',
        ('HOPE_DRIVE','C8'): 'Adams, Alex ~ Learner One',
        ('HOPE_DRIVE','D8'): 'Adams, Alex ~ Learner One',
        ('HOPE_DRIVE','F8'): 'Brown, Blair ~ Learner One',
        ('HOPE_DRIVE','B10'): 'Available, Same ~ ',
        ('NYES','B6'): 'Available, Other ~ ',
        ('ETOWN','C6'): 'Adams, Alex ~ Learner One', # duplicated across Academic Pediatrics; once
        ('NYES','E6'): 'Adams, Alex ~ Learner Two',
        ('WARD A','D6'): 'Chen, Chris ~ Learner Three',
        ('PSHCH_NURSERY','E6'): 'Diaz, Drew ~ Learner Four',
        ('COMPLEX','D16'): 'Adams, Alex ~ Learner Two',
        ('SJR_HOSP','B6'): 'SJR_1 ~ Learner Five',
        ('HOPE_DRIVE','B32'): 'Brown, Blair ~ Learner Two',
    }
    if availability_only:
        values = {k: v.split('~')[0]+'~ ' for k,v in values.items()}
    values.update(changes or {})
    for (site, cell), value in values.items():
        wb[site][cell] = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make_roster(start=date(2026,8,3)):
    return ('legal_name,start_date\n'+''.join(f'Learner {word},{start.isoformat()}\n'
            for word in ('One','Two','Three','Four','Five'))).encode()


def make_qgenda(start=date(2026,8,3)):
    wb=Workbook(); ws=wb.active
    ws.append(['academic general pediatrics hospitalists complex care'])
    tasks=['hope drive am continuity','hope drive pm continuity','hope drive am acute precept',
           'hope drive pm acute precept','etown am continuity','etown pm continuity','nyes rd am continuity',
           'nyes rd pm continuity','nursery weekday 8a-6p','rounder 1 7a-7p','rounder 2 7a-7p',
           'rounder 3 7a-7p','hope drive clinic am','hope drive clinic pm']
    for i in range(28):
        ws.append([(start+timedelta(days=i)).strftime('%B %d, %Y'), ''])
        for j,task in enumerate(tasks):
            ws.append([task, f'Example Provider {j+1}'])
    out=io.BytesIO(); wb.save(out); return out.getvalue()


def canonical(raw):
    """Compare package contents, not ZIP timestamps or document creation dates."""
    if isinstance(raw, str):
        raw=raw.encode()
    if raw[:2] == b'PK':
        with ZipFile(io.BytesIO(raw)) as z:
            return {name: canonical(z.read(name)) for name in sorted(z.namelist())}
    if raw.lstrip().startswith(b'<?xml') or raw.lstrip().startswith(b'<'):
        try:
            root=ET.fromstring(raw)
            for el in root.iter():
                if el.tag in {'{http://purl.org/dc/terms/}created','{http://purl.org/dc/terms/}modified'}:
                    el.text='TIME'
            raw=ET.tostring(root)
        except ET.ParseError:
            pass
    # Only this explanatory help text changes with the mapping's new location.
    return raw.replace(b'PRECEPTOR_EMAIL_MAP in app_sch_2026.py', b'PRECEPTOR_EMAIL_MAP in schedule_app/settings.py')


def run_app(values, *, original=None, secrets=None, state=None, repo=None, evaluation_login=False):
    st.reset(values=values, secrets=secrets, state=state)
    if evaluation_login:
        from schedule_app.services.evaluation_access import _login_callback, P as ACCESS_P
        st.secrets["evaluation_access"] = {"password": "test-only-evaluation-passphrase-2026"}
        st.session_state[ACCESS_P + "password"] = st.secrets["evaluation_access"]["password"]
        _login_callback()
    for value in values.values():
        values_to_seek = value if isinstance(value,list) else [value]
        for upload in values_to_seek:
            if isinstance(upload,Upload):upload.seek(0)
    random.seed(90210)
    with patch('requests.request', (repo or FakeGitHub()).request):
        try:
            runpy.run_path(str(original or ROOT/'app_sch_2026.py'), run_name='__main__')
        except StopRun:
            pass
    return {'downloads':dict(st.downloads), 'messages':list(st.messages),
            'state':dict(st.session_state), 'events':list(st.events), 'widgets':list(st.widget_keys)}


def login_for_test():
    """Explicit authenticated fixture for protected UI/callback functional tests."""
    from schedule_app.services.evaluation_access import _login_callback, P as ACCESS_P
    st.secrets["evaluation_access"] = {"password": "test-only-evaluation-passphrase-2026"}
    st.session_state[ACCESS_P + "password"] = st.secrets["evaluation_access"]["password"]
    _login_callback()
