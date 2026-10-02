"""Optional presentation overrides. Defaults are a strict no-op."""
from docx.shared import Pt
from schedule_app.services.report_wording import active_appearance, report_text


def add_report_message(doc, key):
    text = report_text(key)
    if not text.strip():
        return
    for block in text.replace('\r\n', '\n').replace('\r', '\n').split('\n\n'):
        p = doc.add_paragraph(block)
        p.paragraph_format.keep_with_next = False


def apply_report_appearance(doc, group):
    settings = active_appearance(group)
    if not settings:
        return  # Preserve every original style/run if Admin has made no changes.
    if 'font' in settings:
        for style in doc.styles:
            if hasattr(style, 'font'):
                style.font.name = settings['font']
    if 'body_size' in settings:
        doc.styles['Normal'].font.size = Pt(settings['body_size'])
    if 'title_size' in settings:
        doc.styles['Title'].font.size = Pt(settings['title_size'])

    first_heading = next((p._p for p in doc.paragraphs if p.style and p.style.name in ('Title', 'Heading 1')), None)

    def apply(p, *, table=False, header=False):
        for run in p.runs:
            if 'font' in settings:
                run.font.name = settings['font']
            if header:
                continue  # Page-number fields and header/footer geometry remain intact.
            style = p.style.name if p.style else ''
            if table:
                size = settings.get('table_size')
            elif style == 'Title' or (group in ('instructions', 'opd_changes', 'assignment_summary') and p._p is first_heading):
                size = settings.get('title_size')
            elif style.startswith('Heading') or style == 'Subtitle':
                size = None
            elif run.font.size is not None and run.font.size.pt <= 9:
                size = settings.get('note_size')
            else:
                size = settings.get('body_size')
            if size is not None:
                run.font.size = Pt(size)
    for p in doc.paragraphs:
        apply(p)
    seen = set()
    def tables(items):
        for table in items:
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    for p in cell.paragraphs:
                        apply(p, table=True)
                    tables(cell.tables)
    tables(doc.tables)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                apply(p, header=True)
