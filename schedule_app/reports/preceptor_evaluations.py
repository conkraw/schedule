"""An OASIS evaluation section appended to a matched preceptor's Word report.

Actual question wording is rendered, never technical CSV column identifiers.
Evaluation ratings are kept distinct from scheduled hours and unique learners.
"""
from datetime import date
from docx.shared import Pt, RGBColor
from schedule_app.reports.teaching_tables import teaching_add_work_table


def append_oasis_evaluations(doc, preceptor_name, feedback):
    if feedback is None:
        return
    heading = doc.add_heading("Learner feedback on teaching", level=1)
    heading.paragraph_format.page_break_before = True
    doc.add_paragraph(preceptor_name, style="Subtitle")
    start, end = date.fromisoformat(feedback["start_date"]), date.fromisoformat(feedback["end_date"])
    doc.add_paragraph(f"OASIS Submit Date: {start:%B} {start.day}, {start.year} - {end:%B} {end.day}, {end.year}")
    p = doc.add_paragraph()
    p.add_run("Submitted evaluations: ").bold = True
    p.add_run(f"{feedback['evaluation_count']:,}")
    p.add_run("   |   Username: ").bold = True
    p.add_run(feedback["record_id"])
    p.paragraph_format.keep_with_next = True

    def note(text):
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(5)
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string("526475")
        return p

    note("Evaluation count is the number of submitted forms, not the number of questions or unique students. "
         "This section describes evaluations submitted during these dates; it is not restricted to the students counted in the OPD teaching section.")
    if feedback["educator_name"] != preceptor_name:
        note("Educator name in the selected OASIS summary: " + feedback["educator_name"] + ". Matched by your saved username link.")
    doc.add_heading("Teaching questions", level=2)
    questions = [q for q in feedback["questions"] if not q["duration_category"]]
    if questions:
        teaching_add_work_table(doc, ("Question", "Mean value", "Responses"),
            [(q["question"], q["mean"], f"{q['response_count']:,}") for q in questions],
            widths=(5.2, 0.85, 0.85), number_columns=(1, 2))
    else:
        doc.add_paragraph("No multiple-choice teaching questions were present in this saved summary.")
    note("Means use the exported Multiple Choice Value for each question. Blank/N/A ratings are excluded; "
         "Responses is the number of scored answers for that question. Not scored means no numeric answers were available. "
         "No composite score is calculated.")
    duration = [q for q in feedback["questions"] if q["duration_category"]]
    if duration:
        doc.add_heading("Time with the preceptor", level=2)
        teaching_add_work_table(doc, ("Question", "Mean code", "Responses"),
            [(q["question"], q["mean"], f"{q['response_count']:,}") for q in duration],
            widths=(5.2, 0.85, 0.85), number_columns=(1, 2))
        note("This is a mean of duration-category codes, not weeks, days, or a teaching-quality rating.")
    # A fresh comment page prevents crowded tables and provides a clear separation
    # between question scores and qualitative feedback; long comments flow freely.
    for index, (title, field) in enumerate((("Please indicate this educator's strengths", "strengths_comments"),
                                           ("Areas for Improvement", "areas_for_improvement_comments"))):
        p = doc.add_heading(title, level=2)
        p.paragraph_format.page_break_before = index == 0
        text = feedback[field]
        if not text.strip():
            doc.add_paragraph("No comment recorded in this saved summary.")
        else:
            # Preserve wording, numbering, embedded newlines, and repeat comments
            # from different evaluations. Do not insert HTML or interpret markup.
            for part in text.replace("\r\n", "\n").replace("\r", "\n").split("\n\n"):
                p = doc.add_paragraph(part)
                p.paragraph_format.keep_with_next = False
    note("Comments are reproduced from the selected summary without rewriting. They may contain identifying details; "
         "handle this document as evaluation data.")
    note("Source: " + feedback["summary_filename"] + ". Saved report label: " + feedback["oasis_label"] +
         ". Source file identifier: " + feedback["summary_sha"][:12] + ".")


def append_feedback_unavailable(doc, preceptor_name, reason):
    """Requested feedback must not disappear silently when no link/record exists."""
    heading = doc.add_heading("Learner feedback on teaching", level=2)
    heading.paragraph_format.keep_with_next = True
    paragraph = doc.add_paragraph("Not attached: " + reason.rstrip(".") + ".")
    paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph("This is not a zero evaluation score or proof that no feedback was submitted. "
                      "Check the preceptor username in PTS Matching and the exact-date saved OASIS summary in PTS.")
