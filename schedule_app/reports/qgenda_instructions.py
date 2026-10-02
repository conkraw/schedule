"""Editable QGenda instructions; dates still come from the selected schedule."""
from datetime import datetime, timedelta
from io import BytesIO
from docx import Document
from docx.shared import Pt
from schedule_app.services.report_wording import report_text as rt, with_report_wording, TOKEN, effective_text, ACTIVE
from schedule_app.reports.report_appearance import apply_report_appearance, add_report_message


def _step(doc, template, values):
    p = doc.add_paragraph(style='List Bullet')
    last = 0
    for match in TOKEN.finditer(template):
        p.add_run(template[last:match.start()])
        p.add_run(str(values[match.group(1)])).bold = True
        last = match.end()
    if template[last:]:
        p.add_run(template[last:])


@with_report_wording
def make_qgenda_instructions(start, end):
    doc = Document()
    doc.add_heading(rt('instructions.title'), level=1)
    doc.styles['Normal'].font.size = Pt(8)
    p = doc.add_paragraph()
    p.add_run(rt('instructions.date_prefix'))
    p.add_run(f'{start:%B %d, %Y}').bold = True
    p.add_run(' → ')
    p.add_run(f'{end:%B %d, %Y}').bold = True
    add_report_message(doc, 'instructions.opening')
    intro = rt('instructions.intro')
    if intro:
        doc.add_paragraph(intro)
    values = {'start_date': f'{start:%m/%d/%Y}', 'end_date': f'{end:%m/%d/%Y}'}
    for index in range(1, 5):
        doc.add_heading(rt(f'instructions.report_{index}_title'), level=2)
        template = effective_text(ACTIVE.get(), f'instructions.report_{index}_steps')
        for line in template.replace('\r\n', '\n').replace('\r', '\n').split('\n'):
            if line.strip():
                _step(doc, line, values)
    add_report_message(doc, 'instructions.closing')
    apply_report_appearance(doc, 'instructions')
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
