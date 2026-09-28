"""Shared Word table formatting for teaching reports.

Extracted from the supplied app; this module performs no page rendering on import.
"""

from docx.shared import Pt


def teaching_add_work_table(doc, headers, entries, *, widths, number_columns=(), total=None):
    """A compact, editable table with repeating headers and readable page breaks."""
    from docx.shared import Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    measured = tuple(Inches(value) for value in widths)
    for col, width in zip(table.columns, measured):
        col.width = width
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    def write_row(row, values, *, heading=False, subtotal=False, band=False):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, (cell, value, width) in enumerate(zip(row.cells, values, measured)):
            cell.width = width
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margin = OxmlElement("w:tcMar")
            for edge, amount in (("top", 45), ("bottom", 45), ("left", 90), ("right", 90)):
                element = OxmlElement("w:" + edge)
                element.set(qn("w:w"), str(amount))
                element.set(qn("w:type"), "dxa")
                margin.append(element)
            properties.append(margin)
            if heading or subtotal or band:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "24466B" if heading else "E8EEF5" if subtotal else "F5F7FA")
                properties.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.keep_with_next = heading
            if i in number_columns:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                run.font.name, run.font.size = "Calibri", Pt(10)
                run.bold = heading or subtotal
                if heading:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    write_row(table.rows[0], headers, heading=True)
    for index, values in enumerate(entries):
        write_row(table.add_row(), values, band=index % 2 == 1)
    if total is not None:
        write_row(table.add_row(), total, subtotal=True)
        if len(table.rows) > 2:
            for cell in table.rows[-2].cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    return table
