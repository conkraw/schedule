"""Individual preceptor Word report, by academic year and type of work.

Extracted from the supplied app; this module performs no page rendering on import.
"""

from schedule_app.services.teaching_priority import outpatient_priority_report_note
from schedule_app.services.educational_time import (
    with_educational_time, teaching_time_rows, TIME_DEFINITION, TIME_SCOPE_NOTE,
)
from schedule_app.services.report_diagnostics import checked_report_reach, ReportDataError
from collections import defaultdict
from datetime import date as CalendarDate
from docx import Document
from docx.shared import Pt
from io import BytesIO
from schedule_app.reports.teaching_tables import teaching_add_work_table
from schedule_app.services.opd_archive import OPDArchiveError
from schedule_app.services.reporting_periods import (
    teaching_report_bounds, teaching_report_label, teaching_report_heading, teaching_report_date_text,
)
from schedule_app.services.teaching_analysis import teaching_brief_months
from schedule_app.services.teaching_analysis import teaching_month_label
from schedule_app.services.teaching_analysis import teaching_require_work_type_data
from schedule_app.services.teaching_analysis import teaching_work_type_sort
from schedule_app.settings import TEACHING_HOURS_PER_STUDENT_SHIFT
from schedule_app.settings import TEACHING_WORK_TYPE_REVIEW
from schedule_app.services.learner_reach import (
    learner_reach_rows, reach_percent, reach_totals, require_learner_reach_data,
    REACH_SCOPE_NOTE, PARTICIPATION_SCOPE_NOTE, REACH_DETAIL_TOTAL_NOTE, participating_reach_rows,
)

from schedule_app.services.student_continuity import (
    require_student_continuity_data, student_continuity_counts,
)


def teaching_make_docx(name, monthly, scan, *, oasis_feedback=None, _batch=None):
    """One document per preceptor; academic years and work types stay separate."""
    from docx.shared import Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    teaching_require_work_type_data(scan)
    require_learner_reach_data(scan)
    require_student_continuity_data(scan)
    monthly = [row for row in monthly
               if row["preceptor_name"] == name and int(row["no_of_shifts"]) > 0]
    if not monthly:
        raise OPDArchiveError("No student assignments were found for this preceptor in the selected reporting period; no individual report was generated.")
    from schedule_app.services.teaching_validation import validate_teaching_report
    validate_teaching_report(scan, {row["academic_start_year"] for row in monthly})
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    section.header_distance = section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string("24466B")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8 if style_name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(4)
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            border = properties.find(qn("w:pBdr"))
            if border is not None:
                properties.remove(border)
    subtitle = doc.styles["Subtitle"]
    subtitle.font.name, subtitle.font.size = "Calibri", Pt(12)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor.from_string("526475")
    subtitle.paragraph_format.space_after = Pt(7)
    subtitle.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "PENN STATE  |  PEDIATRIC CLERKSHIP"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("526475")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Third-year student teaching  |  Page ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = f"Preceptor teaching report - {name}"
    doc.core_properties.author = "Pediatric Clerkship"
    doc.core_properties.subject = "Total scheduled availability, educational hours and student continuity"

    def note(text, warning=False):
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("8C3B25" if warning else "526475")
        return paragraph

    grouped = defaultdict(list)
    for row in monthly:
        grouped[row["academic_start_year"]].append(row)
    selected = sorted(grouped)
    if _batch is None:
        overall = teaching_time_rows(scan, selected)
        services = teaching_time_rows(scan, selected, by_work_type=True)
        months = teaching_time_rows(scan, selected, by_work_type=True, monthly=True)
    else:
        from schedule_app.reports.teaching_batch import TeachingReportBatch
        if not isinstance(_batch, TeachingReportBatch):
            raise OPDArchiveError("Invalid report preparation context. Rebuild the reports.")
        overall, services, months = _batch.preceptor_rows(scan, name, selected)
    overall_rows = {row["academic_year"]: row for row in overall if row["preceptor_name"] == name}
    service_rows = [row for row in services if row["preceptor_name"] == name]
    monthly_rows = [row for row in months if row["preceptor_name"] == name]
    for index, year in enumerate(selected):
        label = teaching_report_label(scan, year)
        report_title = doc.add_paragraph("Preceptor teaching report", style="Subtitle")
        report_title.paragraph_format.page_break_before = bool(index)
        doc.add_paragraph(name, style="Title")
        doc.add_heading(teaching_report_heading(scan, year), level=1)
        doc.add_paragraph(teaching_report_date_text(scan, year))
        if name in scan["unresolved_preceptor_labels"]:
            note("Review required: this is a site, slot, or combined provider label, not a verified individual preceptor.", warning=True)
        reach = overall_rows.get(label)
        if reach is None:
            raise ReportDataError("No clinical metrics were found for this included preceptor.",
                                  report="Individual preceptor report", preceptor_name=name,
                                  academic_year=label)
        doc.add_heading("Teaching time", level=2)
        teaching_add_work_table(doc, ("Measure", "Result"), [
            ("Total scheduled availability", f"{reach['total_scheduled_availability_hours']:,} hours"),
            ("Educational hours", f"{reach['educational_hours']:,} hours"),
            ("Learner Reach", reach_percent(reach["learner_reach_pct"])),
        ], widths=(4.9, 2.0), number_columns=(1,))
        note(f"{reach['teaching_shifts']:,} of {reach['scheduled_shifts']:,} scheduled AM/PM shifts included at least one student.")
        note(TIME_DEFINITION)
        doc.add_heading("Student continuity", level=2)
        p = doc.add_paragraph()
        p.add_run("Unique students assigned: ").bold = True
        p.add_run(f"{reach['unique_students']:,}")
        p.add_run("   |   Students assigned on 3+ days: ").bold = True
        p.add_run(f"{reach['unique_students_3plus_days']:,}")
        p.paragraph_format.keep_with_next = True
        note("Students are counted once across all work types in this period. Three days means three distinct dates; "
             "AM and PM on the same date count as one day. These counts do not multiply educational hours.")
        reach_types = [row for row in service_rows if row["academic_year"] == label]
        doc.add_heading("By clinical experience", level=2)
        titles = ("Clinical experience", "Total scheduled\navailability (hours)", "Educational\nhours", "Learner Reach")
        values = [(row["work_type"], f"{row['total_scheduled_availability_hours']:,}",
                   f"{row['educational_hours']:,}", reach_percent(row["learner_reach_pct"]))
                  for row in reach_types]
        teaching_add_work_table(doc, titles, values,
            widths=(2.45, 1.75, 1.4, 1.3), number_columns=(1, 2, 3),
            total=("Overall", f"{reach['total_scheduled_availability_hours']:,}",
                   f"{reach['educational_hours']:,}", reach_percent(reach["learner_reach_pct"])))
        if sum(row["total_scheduled_availability_hours"] for row in reach_types) != reach["total_scheduled_availability_hours"]:
            note(REACH_DETAIL_TOTAL_NOTE)
        month_values = [(row["work_type"] + " — " + teaching_month_label(CalendarDate.fromisoformat(row["month"])),
                         f"{row['total_scheduled_availability_hours']:,}", f"{row['educational_hours']:,}",
                         reach_percent(row["learner_reach_pct"]))
                        for row in monthly_rows if row["academic_year"] == label]
        # Put longer monthly breakdowns on a fresh page so their notes do not
        # spill alone onto a near-empty final teaching page.
        monthly_heading = doc.add_heading("Monthly detail", level=2)
        monthly_heading.paragraph_format.page_break_before = len(month_values) > 6
        teaching_add_work_table(doc,
            ("Clinical experience / month", "Total scheduled\navailability (hours)", "Educational\nhours", "Learner Reach"),
            month_values, widths=(2.45, 1.75, 1.4, 1.3), number_columns=(1, 2, 3))
        note("Only clinical experiences with student assignments in the selected period are listed. "
             "For those experiences, months without students remain included in availability.")
        priority_note = outpatient_priority_report_note(scan, [year], preceptor_name=name)
        if priority_note:
            note(priority_note)
        note(TIME_SCOPE_NOTE)
        note("Source: current encrypted OPDs. "
             f"Snapshot: {scan['commit'][:12]}; retrieved: {scan['generated_at']}. Student names omitted.")
        from schedule_app.reports.assessment_completion import append_individual_completion
        append_individual_completion(doc, scan, name, year)
        if oasis_feedback is not None:
            from schedule_app.services.teaching_evaluations import feedback_for_preceptor
            from schedule_app.reports.preceptor_evaluations import append_oasis_evaluations
            append_oasis_evaluations(doc, name, feedback_for_preceptor(oasis_feedback, scan, name, year))
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
