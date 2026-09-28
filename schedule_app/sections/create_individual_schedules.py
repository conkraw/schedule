"""Sidebar section: Create Individual Schedules.

Extracted from the supplied app; this module performs no page rendering on import.
"""

from collections import defaultdict
from io import BytesIO
from openpyxl import load_workbook
from schedule_app.services.primary_preceptors import build_preceptor_assignment_report
from schedule_app.services.primary_preceptors import build_preceptor_report_workbook
from schedule_app.services.workbook_copy import copy_sheet_to_new_wb
from schedule_app.settings import INDIVIDUAL_OUTPUT_STATE_KEYS
from schedule_app.settings import INDIVIDUAL_REPORT_SCHEMA_VERSION
from zipfile import ZIP_DEFLATED
from zipfile import ZipFile
import hashlib
import pandas as pd
import re
import streamlit as st


def render():
    """Render the Create Individual Schedules sidebar section."""
    st.subheader("Individual Schedule Creator")

    # Edit email mappings in schedule_app/settings.py.
    # Existing session keys and schema-version checks are preserved below.

    if (
        st.session_state.get("individual_report_schema_version")
        != INDIVIDUAL_REPORT_SCHEMA_VERSION
    ):
        for state_key in INDIVIDUAL_OUTPUT_STATE_KEYS:
            st.session_state.pop(state_key, None)
        st.session_state["individual_report_schema_version"] = (
            INDIVIDUAL_REPORT_SCHEMA_VERSION
        )

    uploaded = st.file_uploader(
        "Upload the master Excel (.xlsx) with one tab per person",
        type=["xlsx"],
        key="individual_schedule_master",
    )


    if uploaded is not None:
        # Clear previously generated output when a different source file is uploaded.
        source_signature = (uploaded.name, hashlib.sha256(uploaded.getvalue()).hexdigest())
        if st.session_state.get("individual_schedule_source") != source_signature:
            st.session_state["individual_schedule_source"] = source_signature
            st.session_state.pop("individual_schedule_zip", None)
            st.session_state.pop("individual_preceptor_report", None)
            st.session_state.pop("individual_preceptor_preview", None)
            st.session_state.pop("individual_missing_emails", None)

        # Keep formulas/formatting -> data_only=False
        wb = load_workbook(uploaded, data_only=False)
        st.write(f"Found **{len(wb.sheetnames)}** tabs.")
        st.caption(
            "The preceptor report includes only HOPE_DRIVE, NYES, and ETOWN. "
            "Each student/week with focus-site assignments receives one primary preceptor. The app prefers "
            ">=3 sessions; fallback and repeated primary assignments are flagged. "
            "Fragmented = <3 sessions."
        )

        if st.button("Build individual schedules + preceptor report"):
            report_df = build_preceptor_assignment_report(wb)
            report_buf, missing_emails = build_preceptor_report_workbook(report_df)

            zip_buf = BytesIO()
            with ZipFile(zip_buf, mode="w", compression=ZIP_DEFLATED) as zf:
                used_names = defaultdict(int)

                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]

                    # Skip truly empty sheets (no cells with value)
                    has_any_value = any(
                        cell.value is not None
                        for row in ws.iter_rows()
                        for cell in row
                    )
                    if not has_any_value:
                        continue

                    # Safe file name
                    base = re.sub(r"[^A-Za-z0-9._-]+", "_", sheet_name).strip("_") or "sheet"
                    used_names[base] += 1
                    safe_name = (
                        base
                        if used_names[base] == 1
                        else f"{base}_{used_names[base]}"
                    )

                    # Copy this sheet into its own new workbook (preserving formatting)
                    out_buf = copy_sheet_to_new_wb(ws)
                    zf.writestr(f"{safe_name}.xlsx", out_buf.getvalue())

                # Include the new report in the same ZIP.
                zf.writestr(
                    "Preceptor_Assignment_Report.xlsx", report_buf.getvalue()
                )

            zip_buf.seek(0)
            st.session_state["individual_schedule_zip"] = zip_buf.getvalue()
            st.session_state["individual_preceptor_report"] = report_buf.getvalue()
            st.session_state["individual_preceptor_preview"] = report_df
            st.session_state["individual_missing_emails"] = missing_emails

        report_preview = st.session_state.get("individual_preceptor_preview")

        # Defensive protection for stale output created by an older app version.
        # A prior preview may be a valid DataFrame but lack newly added columns.
        required_preview_columns = {
            "primary_preceptor",
            "primary_preceptor_flag",
            "primary_preceptor_flag_reason",
        }
        preview_is_current = (
            isinstance(report_preview, pd.DataFrame)
            and required_preview_columns.issubset(report_preview.columns)
        )
        if report_preview is not None and not preview_is_current:
            for state_key in INDIVIDUAL_OUTPUT_STATE_KEYS:
                st.session_state.pop(state_key, None)
            report_preview = None
            st.info(
                "The preceptor report format was updated. Please click "
                "'Build individual schedules + preceptor report' to regenerate both outputs."
            )

        if report_preview is not None:
            st.markdown("**Preceptor assignment report preview**")
            if report_preview.empty:
                st.warning(
                    "No HOPE_DRIVE, NYES, or ETOWN assignments were found in the uploaded workbook."
                )
            else:
                st.dataframe(report_preview, use_container_width=True, hide_index=True)

            flagged_primary_rows = report_preview.loc[
                report_preview["primary_preceptor"].eq("YES")
                & report_preview["primary_preceptor_flag"].eq("YES")
            ]
            if not flagged_primary_rows.empty:
                st.warning(
                    f"{len(flagged_primary_rows)} primary assignment(s) require review. "
                    "See primary_preceptor_flag_reason in the preview or Excel report."
                )
            elif not report_preview.empty:
                st.success("Every primary assignment met the preferred criteria without reuse flags.")

            missing_emails = st.session_state.get("individual_missing_emails", [])
            if missing_emails:
                st.warning(
                    f"{len(missing_emails)} preceptor(s) are missing from "
                    "PRECEPTOR_EMAIL_MAP. Their email cells are blank, and their names "
                    "are listed on the 'Missing Emails' tab."
                )
            elif not report_preview.empty:
                st.success("All preceptors in this report have a mapped email address.")

        if st.session_state.get("individual_preceptor_report") is not None:
            st.download_button(
                label="Download preceptor assignment report (Excel)",
                data=st.session_state["individual_preceptor_report"],
                file_name="Preceptor_Assignment_Report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        if st.session_state.get("individual_schedule_zip") is not None:
            st.download_button(
                label="Download individual schedules + report (ZIP)",
                data=st.session_state["individual_schedule_zip"],
                file_name="individual_schedules_with_preceptor_report.zip",
                mime="application/zip",
            )
