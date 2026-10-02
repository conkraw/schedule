"""Password-protected report wording editor. No Python/JSON editing required."""
from copy import deepcopy
from io import BytesIO
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from schedule_app.services.evaluation_access import (
    require_evaluation_access, evaluation_access_is_valid, lock_evaluation_records,
)
from schedule_app.services.opd_archive import GitHubOPDArchive, get_opd_archive_config, OPDArchiveError
from schedule_app.services.report_wording import (
    GitHubReportWording, report_catalog, empty_wording, wording_data, effective_text,
    validate_field, validate_wording, report_wording_context, render_text,
    FONT_CHOICES, STYLE_GROUPS, STYLE_LIMITS,
)
from schedule_app.sections.report_wording_controls import accept_saved_report_wording
from schedule_app.reports.report_appearance import apply_report_appearance

P = '_admin_report_'


def _preview(snapshot, group, key):
    """A wording proof, not invented evaluation results or a live report."""
    field = report_catalog()['fields'][key]
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = section.bottom_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    doc.add_heading('Report wording preview', 0)
    doc.add_paragraph(report_catalog()['groups'][group], style='Subtitle')
    doc.add_paragraph('This is a wording proof, not a calculated report. Sample date and shift values below are illustrative.')
    doc.add_heading(field['label'], 1)
    values = {'start_date': '03/01/2026', 'end_date': '03/01/2027', 'minimum_shifts': 3, 'shift_word': 'shifts', 'start_long': 'March 01, 2026', 'end_long': 'March 01, 2027'}
    text = render_text(effective_text(snapshot, key), values)
    doc.add_paragraph(text or '(No additional wording.)')
    if group in STYLE_GROUPS:
        with report_wording_context(snapshot):
            apply_report_appearance(doc, group)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _save(store, snapshot, data, message):
    # Recheck authorization immediately before the write, not only at page entry.
    if not evaluation_access_is_valid(touch=True):
        lock_evaluation_records()
        st.warning('Access expired. Unlock Admin before saving report wording.')
        return False
    try:
        saved = store.save(data, expected=snapshot)
    except OPDArchiveError as exc:
        st.error(str(exc))
        return False
    st.session_state[P + 'snapshot'] = saved
    st.session_state[P + 'notice'] = message
    st.session_state[P + 'revision'] = st.session_state.get(P + 'revision', 0) + 1
    st.session_state.pop(P + 'preview', None)
    accept_saved_report_wording(saved)
    st.rerun()
    return True


def render():
    if not require_evaluation_access(section_name='Admin', lock_key='evaluation_lock_admin'):
        return
    st.subheader('Admin')
    st.write('Edit report instructions and wording here—no Python or JSON editing. Saved changes apply to newly generated reports.')
    st.caption('Current wording remains the default. Calculations, report dates, usernames, CSV columns, question wording from OASIS '
               'and evaluation comments are not edited here. Changes are shared by everyone using this archive.')
    try:
        client = GitHubOPDArchive(get_opd_archive_config())
        store = GitHubReportWording(client)
    except OPDArchiveError as exc:
        st.error(str(exc))
        return
    scope = client.config.signature()
    if st.session_state.get(P + 'scope') != scope:
        for key in list(st.session_state):
            if str(key).startswith(P):
                st.session_state.pop(key, None)
        st.session_state[P + 'scope'] = scope
    if st.button('Reload saved wording from GitHub', key=P + 'reload'):
        st.session_state.pop(P + 'snapshot', None)
        st.session_state.pop(P + 'preview', None)
        st.session_state[P + 'revision'] = st.session_state.get(P + 'revision', 0) + 1
    if P + 'snapshot' not in st.session_state:
        try:
            st.session_state[P + 'snapshot'] = store.load()
        except OPDArchiveError as exc:
            st.error('Saved wording could not be loaded. ' + str(exc))
            return
    snapshot = st.session_state[P + 'snapshot']
    notice = st.session_state.pop(P + 'notice', None)
    if notice:
        st.success(notice)
    else:
        st.caption('Using saved custom wording.' if snapshot['texts'] or snapshot['styles']
                   else 'No custom wording yet. All existing reports are unchanged.')
    catalog = report_catalog()
    group = st.selectbox('Report or report component', list(catalog['groups']),
                          format_func=lambda g: catalog['groups'][g], key=P + 'group')
    keys = [k for k, v in catalog['fields'].items() if v['group'] == group]
    revision = st.session_state.get(P + 'revision', 0)
    key = st.selectbox('Part to edit', keys, format_func=lambda k: catalog['fields'][k]['label'],
                        key=P + f'field_{group}')
    field = catalog['fields'][key]
    st.caption('Edit plain text. Use a blank line for a new paragraph. Changes are not saved until you click Save.')
    if field['tokens']:
        st.info('Keep these automatic placeholders: ' + ', '.join('{' + token + '}' for token in field['tokens']) +
                '. The actual dates or minimum shifts are filled in when a report is created.')
    with st.form(P + f'edit_form_{key}_{revision}'):
        text = st.text_area('Report wording', value=effective_text(snapshot, key), height=240,
                             max_chars=16000, key=P + f'text_{key}_{revision}')
        save = st.form_submit_button('Save wording to GitHub')
        preview = st.form_submit_button('Preview this edit (Word)')
    if save or preview:
        try:
            validate_field(key, text)
            candidate = wording_data(snapshot)
            candidate['texts'][key] = text
            validate_wording(candidate)
            if save:
                _save(store, snapshot, candidate, 'Report wording saved, encrypted, and verified. Generate a new report to use it.')
            else:
                st.session_state[P + 'preview'] = {'group': group, 'key': key, 'bytes': _preview(candidate, group, key)}
        except OPDArchiveError as exc:
            st.error(str(exc))
    proof = st.session_state.get(P + 'preview')
    if proof and proof['group'] == group and proof['key'] == key:
        st.download_button('Download wording preview (Word)', data=proof['bytes'],
                           file_name='Report_Wording_Preview.docx',
                           mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                           key=P + 'preview_download')
        st.caption('Preview edits are not saved. Use Save wording to GitHub to apply them. For a complete layout check, generate the usual report after saving.')
    with st.expander('Original wording / restore this part (optional)'):
        st.text(field['default'] or '(No additional wording by default.)')
        confirm = st.checkbox('Restore this part to its original wording', key=P + f'reset_field_confirm_{key}')
        if st.button('Restore this part', disabled=not confirm, key=P + f'reset_field_{key}'):
            candidate = wording_data(snapshot)
            candidate['texts'].pop(key, None)
            _save(store, snapshot, candidate, 'This part was restored to its original wording.')
    if group in STYLE_GROUPS:
        with st.expander('Word appearance (optional)'):
            st.caption('Keep existing preserves the current layout. Larger fonts or longer notes can add pages. '
                       'Tables, page width, numeric formats and automatic page numbers remain intact.')
            style = snapshot['styles'].get(group, {})
            with st.form(P + f'appearance_{group}_{revision}'):
                fonts = ['Keep existing'] + list(FONT_CHOICES)
                font = st.selectbox('Font', fonts, index=fonts.index(style.get('font', 'Keep existing')),
                                    key=P + f'font_{group}_{revision}')
                selected = {}
                for name, bounds in STYLE_LIMITS.items():
                    choices = ['Keep existing'] + list(range(bounds[0], bounds[1] + 1))
                    # A prior supported fractional size remains selectable.
                    if name in style and style[name] not in choices:
                        choices.append(style[name])
                    value = st.selectbox(name.replace('_', ' ').capitalize() + ' (points)', choices,
                                         index=choices.index(style.get(name, 'Keep existing')),
                                         key=P + f'{name}_{group}_{revision}')
                    if value != 'Keep existing':
                        selected[name] = value
                save_appearance = st.form_submit_button('Save appearance to GitHub')
            if save_appearance:
                candidate = wording_data(snapshot)
                if font != 'Keep existing':
                    selected['font'] = font
                candidate['styles'][group] = selected
                _save(store, snapshot, candidate, 'Report appearance saved, encrypted, and verified.')
    with st.expander('Restore this report or all report settings (optional)'):
        option = st.selectbox('Restore scope', ['Selected report/component', 'All report wording and appearance'],
                              key=P + 'restore_scope')
        confirm = st.checkbox('I want to remove the saved customizations for this scope', key=P + 'restore_confirm')
        if st.button('Restore original report settings', disabled=not confirm, key=P + 'restore'):
            candidate = wording_data(snapshot)
            if option == 'All report wording and appearance':
                candidate = empty_wording()
            else:
                candidate['texts'] = {k: v for k, v in candidate['texts'].items() if k not in keys}
                candidate['styles'].pop(group, None)
            _save(store, snapshot, candidate, 'Original report settings restored. No source files or mappings were deleted.')
    st.caption('Saved separately as ' + store.path + '. No OPDs, evaluations, identity links or calculations are changed. '
               'General scheduling instructions can appear in password-free parts of the app: do not put student information, comments or credentials in these settings. '
               'Restoring defaults changes the current settings only, not Git history or reports already downloaded.')
