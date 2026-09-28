"""Sidebar section: Instructions.

Extracted from the supplied app; this module performs no page rendering on import.
"""

from datetime import datetime
from datetime import timedelta
from docx import Document
from docx.shared import Pt
import io
import streamlit as st


def render():
    """Render the Instructions sidebar section."""
    d = st.text_input('Start date (m/d/yyyy)')
    if d:
        try:
            s = datetime.strptime(d, '%m/%d/%Y')
            e = s + timedelta(days=34)
            st.write(f"{s:%B %d, %Y} → {e:%B %d, %Y}")
    
            st.write('Please go to https://login.qgenda.com/')
    
            # Display on‐screen instructions
            st.markdown(f"""Download four files and create reports based on **{s:%B %d, %Y}** → **{e:%B %d, %Y}**.""")
            st.write("Download instructions here:")


        except ValueError:
            st.error('Invalid format – use m/d/yyyy (e.g. 7/6/2021)')

        # --- Generate a Word document with the same instructions ---
        doc = Document()
        doc.add_heading('Qgenda Report Instructions', level=1)
        doc.styles['Normal'].font.size = Pt(8)
        
        # Bold the date range
        p = doc.add_paragraph()
        p.add_run('Date range: ')
        run_start = p.add_run(f'{s:%B %d, %Y}')
        run_start.bold = True
        p.add_run(' → ')
        run_end = p.add_run(f'{e:%B %d, %Y}')
        run_end.bold = True
        
        
        doc.add_paragraph('1. Go to https://login.qgenda.com/')
        
        # Helper to add each report block
        def add_report(title, steps):
            doc.add_heading(title, level=2)
            for step in steps:
                # If the step contains dates, bold them
                if 'Enter Start Date:' in step:
                    p = doc.add_paragraph(style='List Bullet')
                    prefix, dates = step.split(':', 1)
                    p.add_run(prefix + ': ')
                    # split the two dates on " and End Date:"
                    start_part, end_part = dates.strip().split(' and End Date:')
                    r1 = p.add_run(start_part.strip())
                    r1.bold = True
                    p.add_run(' and End Date: ')
                    r2 = p.add_run(end_part.strip())
                    r2.bold = True
                else:
                    doc.add_paragraph(step, style='List Bullet')
        
        add_report(
            'Report 1 – Penn State Health Hershey Medical Center - Academic General Pediatrics',
            [
                'Click Penn State Health Hershey Medical Center - Academic General Pediatrics → Schedule → Reports',
                'Set Report Type to Calendar by Task',
                'Set Format to Excel',
                f'Enter Start Date: {s:%m/%d/%Y} and End Date: {e:%m/%d/%Y}',
                'Ensure Calendar starts on Monday',
                'Show Staff by Last Name, First Name',
                'Show Tasks by Short Name',
                'Click Run Report'
            ]
        )
        add_report(
            "Report 2 – Penn State Health Children's Hospital – Hospitalists",
            [
                'Click Penn State Health Children\'s Hospital → Schedule → Reports',
                'Set Report Type to Calendar by Task',
                'Set Format to Excel',
                f'Enter Start Date: {s:%m/%d/%Y} and End Date: {e:%m/%d/%Y}',
                'Ensure Calendar starts on Monday',
                'Show Staff by Last Name, First Name',
                'Show Tasks by Long Name',
                'Click Run Report'
            ]
        )
        add_report(
            'Report 3 – Department of Pediatrics (Admin - Adolescent Med)',
            [
                'Click Department of Pediatrics → Schedule → Reports',
                'Select Admin - Adolescent Med in top-right corner',
                'Set Report Type to Calendar by Task',
                'Set Format to Excel',
                f'Enter Start Date: {s:%m/%d/%Y} and End Date: {e:%m/%d/%Y}',
                'Ensure Calendar starts on Monday',
                'Show Staff by Last Name, First Name',
                'Show Tasks by Long Name',
                'Click Run Report'
            ]
        )
        add_report(
            'Report 4 – Department of Pediatrics (Complex Care)',
            [
                'Click Department of Pediatrics → Schedule → Reports',
                'Select Complex Care in top-right corner',
                'Set Report Type to Calendar by Task',
                'Set Format to Excel',
                f'Enter Start Date: {s:%m/%d/%Y} and End Date: {e:%m/%d/%Y}',
                'Ensure Calendar starts on Monday',
                'Show Staff by Last Name, First Name',
                'Show Tasks by Long Name',
                'Click Run Report'
            ]
        )


        # Save to bytes and offer download
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        st.download_button(label="📄 Download Instructions (Word)",data=buf.getvalue(),file_name="Qgenda_Report_Instructions.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
