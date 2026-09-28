"""MD/PA date parsing, booking maps and annotated workbook exports.

Extracted from the supplied app; this module performs no page rendering on import.
"""

from collections import defaultdict
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border
from openpyxl.styles import Color
from openpyxl.styles import Font
from openpyxl.styles import Side
import pandas as pd
import streamlit as st


# -----------------------------
# Config & Constants
# -----------------------------
DAYS = ['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday','Sunday']
DEFAULT_FOCUS = ['HOPE_DRIVE','NYES','ETOWN','LANCASTER']

TRUST_ONLY_AM_PM = True          # Only parse rows with Column A starting AM/PM
REQUIRE_VALID_DATE = True        # Bookings require a valid date anchor


@st.cache_data(show_spinner=False)
def read_sheet_names(file):
    try:
        xl = pd.ExcelFile(file)
        return xl.sheet_names
    except Exception as e:
        st.error(f"Failed to read sheet names: {e}")
        return []


@st.cache_data(show_spinner=False)
def load_sheet(file, sheet_name):
    return pd.read_excel(file, sheet_name=sheet_name, header=None)


def _try_parse_date(x):
    try:
        d = pd.to_datetime(x, errors='coerce')
        if pd.notna(d):
            return d
        return None
    except Exception:
        return None


def find_week_headers(df: pd.DataFrame):
    """
    Robust week header finder.
    - Find 'Monday' (case-insensitive) in Column B.
    - Within next 10 rows, pick first row where B..H has >=2 parseable dates -> dates row.
    - monday_date is B's parsed date; if missing, infer from any parsed day in that row.
    Returns [(monday_row, date_row, monday_date)]
    """
    col = 1
    s = df.iloc[:, col].astype(str).str.strip().str.lower()
    day_rows = df.index[s.eq('monday')].tolist()

    headers = []
    for dr in day_rows:
        date_r, monday_date = None, None
        for look_ahead in range(1, 11):
            r = dr + look_ahead
            if r >= len(df):
                break
            parsed_dates = {}
            for i in range(7):
                c = 1 + i
                if c >= df.shape[1]:
                    continue
                parsed = _try_parse_date(df.iat[r, c])
                if parsed is not None:
                    parsed_dates[i] = parsed
            if len(parsed_dates) >= 2:
                date_r = r
                if 0 in parsed_dates:
                    monday_date = parsed_dates[0].date()
                else:
                    i0 = sorted(parsed_dates.keys())[0]
                    monday_date = (parsed_dates[i0] - pd.Timedelta(days=i0)).date()
                break
        headers.append((dr, date_r, monday_date))
    return headers


def row_to_week_monday(row_idx: int, headers):
    prev = [h for h in headers if h[0] <= row_idx]
    if not prev:
        return None
    prev.sort(key=lambda x: x[0])
    return prev[-1][2]


def detect_am_pm_runs(df: pd.DataFrame, start_row: int = 0):
    """Scan Col A for AM/PM rows and group consecutive runs."""
    runs, current, run_start, prev_idx = [], None, None, None
    for idx in range(start_row, len(df)):
        label = None
        raw = df.iat[idx, 0]
        if isinstance(raw, str):
            ru = raw.strip().upper()
            if ru.startswith('AM'):
                label = 'AM'
            elif ru.startswith('PM'):
                label = 'PM'
        if TRUST_ONLY_AM_PM and label is None:
            continue
        if label is None:
            continue
        if current is None:
            current, run_start = label, idx
        elif label != current or (prev_idx is not None and idx != prev_idx + 1):
            runs.append((current, run_start, prev_idx))
            current, run_start = label, idx
        prev_idx = idx
    if current is not None:
        runs.append((current, run_start, prev_idx))
    return runs


def build_maps_and_roster(df: pd.DataFrame):
    """
    Returns:
      mapping_by_week: {(monday_date, period, day, preceptor) -> {'student','cell','date'}}
                       (only when a real student exists)
      index_by_date:   {(date, period, preceptor) -> {'student','cell','day'}}  # for date-based matching
      roster_week:     {(monday_date, period) -> set(preceptors)}               # week-level roster
      day_roster:      {(monday_date, period, day) -> set(preceptors)}          # day-level roster
      week_dates:      {monday_date -> {day -> date}}
      occupied:        set((monday_date, period, day, preceptor)) even w/o student
      diag_weeks:      diagnostics list
    """
    import re, unicodedata

    def is_placeholder_preceptor(text: str) -> bool:
        if not text: return True
        t = text.strip().upper()
        EXCLUDE_PREFIXES = [
            'CLOSED','CLOSE','BLOCK','VACATION','ADMIN','MEETING','NO CLINIC',
            'CLINIC CANCELLED','CANCELLED','HOLIDAY','OFF','PTO','SICK',
            'NOTE','NOTES','REFERENCE','INFO','FYI','ORIENTATION'
        ]
        return any(t.startswith(pfx) for pfx in EXCLUDE_PREFIXES)

    def _norm(s: str) -> str:
        s = unicodedata.normalize("NFKC", s)
        s = s.replace("\u00A0", " ")                    # NBSP -> space
        s = s.replace("\u2013", "-").replace("\u2014", "-")  # en/em dash -> '-'
        s = re.sub(r"\s+", " ", s.strip())
        return s

    def parse_cell(val: str):
        """Robust split on '~'. RHS counts as student only if alphanumeric & not a placeholder."""
        if not isinstance(val, str): return None, None
        raw = _norm(val)
        if "~" not in raw:
            pre = _norm(raw)
            return (pre if pre else None), None
        pre, rhs = re.split(r"\s*~\s*", raw, maxsplit=1)
        pre = _norm(pre)
        rhs = _norm(rhs)

        BLANK_TOKENS = {"", "nan", "n/a", "na", "-", "--", "—", "none", "null"}
        PLACEHOLDER_HINTS = {"note", "notes", "ref", "reference", "info", "fyi"}

        rhs_l = rhs.lower()
        # treat as empty unless it has at least one letter/number and is not a placeholder
        if (rhs_l in BLANK_TOKENS) or (not re.search(r"[a-z0-9]", rhs_l)) or any(h in rhs_l for h in PLACEHOLDER_HINTS):
            rhs = None
        return (pre if pre else None), rhs

    headers = find_week_headers(df)
    runs = detect_am_pm_runs(df, start_row=0)

    mapping_by_week = {}
    index_by_date = {}
    roster_week = defaultdict(set)
    day_roster = defaultdict(set)
    week_dates = defaultdict(dict)
    occupied = set()
    diag_weeks = []

    # Build date anchors with fallback
    for (day_row, date_row, monday_date) in headers:
        if monday_date is None or date_row is None:
            continue
        inferred_days = []
        for i, day in enumerate(DAYS):
            col_idx = 1 + i  # B..H
            val = df.iat[date_row, col_idx] if col_idx < df.shape[1] else None
            parsed = _try_parse_date(val)
            if parsed is not None:
                week_dates[monday_date][day] = parsed.date()
            else:
                # fallback Monday + i days
                try:
                    fallback = (pd.to_datetime(monday_date) + pd.Timedelta(days=i)).date()
                    week_dates[monday_date][day] = fallback
                    inferred_days.append(day)
                except Exception:
                    week_dates[monday_date][day] = None
                    inferred_days.append(day)
        diag_weeks.append({
            'monday_date': monday_date,
            'date_row': date_row,
            'inferred_days': inferred_days
        })

    # Parse inside AM/PM runs
    for period, rstart, rend in runs:
        monday_date = row_to_week_monday(rstart, headers)
        if monday_date is None:
            continue
        for col_idx, day in enumerate(DAYS, start=1):
            date_anchor = week_dates[monday_date].get(day)
            for row in range(rstart, rend+1):
                if col_idx >= df.shape[1]:
                    continue
                val = df.iat[row, col_idx]
                if pd.isna(val) or not isinstance(val, str):
                    continue
                pre, stu = parse_cell(val)
                if not pre or is_placeholder_preceptor(pre):
                    continue

                # Present in week & specific day
                roster_week[(monday_date, period)].add(pre)
                day_roster[(monday_date, period, day)].add(pre)
                occupied.add((monday_date, period, day, pre))  # presence marker

                # Keep booking only if real student + valid date
                if stu is None:
                    continue
                if REQUIRE_VALID_DATE and _try_parse_date(date_anchor) is None:
                    continue

                cell = f"{chr(ord('A')+col_idx)}{row+1}"
                wk_key = (monday_date, period, day, pre)
                mapping_by_week.setdefault(wk_key, {'student': stu, 'cell': cell, 'date': date_anchor})
                # date-based index for cross-file matching
                dt_key = (pd.to_datetime(date_anchor).date(), period, pre)
                # prefer first seen student for stability
                index_by_date.setdefault(dt_key, {'student': stu, 'cell': cell, 'day': day})

    return mapping_by_week, index_by_date, roster_week, day_roster, week_dates, occupied, diag_weeks


def _annot_make_copy(uploaded_file, other_idx_by_site: dict, selected_sheets: list) -> bytes:
    """
    Annotate: red font (if visible), THICK RED BORDER, and a small note so conflicts
    are obvious even when Conditional Formatting overrides font color.
    """
    raw = uploaded_file.getvalue()
    wb = load_workbook(BytesIO(raw))

    # --- helpers matching your main parser ---
    import re, unicodedata, pandas as pd
    def _norm(s: str) -> str:
        s = unicodedata.normalize("NFKC", s).replace("\u00A0", " ")
        s = s.replace("\u2013", "-").replace("\u2014", "-")
        return re.sub(r"\s+", " ", s.strip())

    def _parse_cell(val: str):
        if not isinstance(val, str): return None, None
        raw = _norm(val)
        if "~" not in raw:
            pre = _norm(raw); return (pre if pre else None), None
        pre, rhs = re.split(r"\s*~\s*", raw, maxsplit=1)
        pre = _norm(pre); rhs = _norm(rhs)
        if rhs.lower() in {"", "nan", "n/a", "na", "-", "--", "—", "none", "null"}:
            rhs = None
        return (pre if pre else None), rhs

    def _is_placeholder_preceptor(text: str) -> bool:
        if not text: return True
        t = text.strip().upper()
        return any(t.startswith(pfx) for pfx in [
            'CLOSED','CLOSE','BLOCK','VACATION','ADMIN','MEETING','NO CLINIC',
            'CLINIC CANCELLED','CANCELLED','HOLIDAY','OFF','PTO','SICK',
            'NOTE','NOTES','REFERENCE','INFO','FYI','ORIENTATION'
        ])

    # opaque ARGB
    OPAQUE_RED = Color(rgb="FFFF0000")
    RED_SIDE   = Side(style="thick", color="FFFF0000")
    RED_BORDER = Border(left=RED_SIDE, right=RED_SIDE, top=RED_SIDE, bottom=RED_SIDE)

    for sheet in selected_sheets:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]

        # rebuild date map from the uploaded file (aligns weeks/days to dates)
        df = load_sheet(uploaded_file, sheet)
        headers = find_week_headers(df)
        runs = detect_am_pm_runs(df, start_row=0)

        week_dates = {}
        for (_day_row, date_row, monday_date) in headers:
            if monday_date is None or date_row is None:
                continue
            week_dates.setdefault(monday_date, {})
            for i, day in enumerate(DAYS):
                c = 1 + i  # B..H
                if c >= df.shape[1]: continue
                parsed = pd.to_datetime(df.iat[date_row, c], errors='coerce')
                if pd.notna(parsed):
                    week_dates[monday_date][day] = parsed.date()

        other_idx = other_idx_by_site.get(sheet, {})  # keys: (date, period, preceptor)

        for period, rstart, rend in runs:
            monday_date = row_to_week_monday(rstart, headers)
            if monday_date is None:
                continue
            for c_idx, day in enumerate(DAYS, start=1):  # B..H
                dt = week_dates.get(monday_date, {}).get(day)
                if dt is None:
                    continue
                for row in range(rstart, rend+1):
                    if c_idx >= df.shape[1]: continue
                    val = df.iat[row, c_idx]
                    if pd.isna(val) or not isinstance(val, str): continue
                    pre, _stu = _parse_cell(val)
                    if not pre or _is_placeholder_preceptor(pre): continue

                    if (dt, period, pre) in other_idx:
                        addr = f"{chr(ord('A')+c_idx)}{row+1}"
                        cell = ws[addr]

                        # Try to ensure red font (CF may still override)
                        f = cell.font or Font()
                        try:
                            cell.font = f.copy(color="FFFF0000")
                        except Exception:
                            cell.font = Font(
                                name=f.name, size=f.size or 11, bold=f.bold,
                                italic=f.italic, underline=f.underline, color=OPAQUE_RED
                            )

                        # Add thick red border (highly visible even with CF)
                        cell.border = RED_BORDER

                        # Add a small note/comment (red triangle)
                        if cell.comment is None:
                            txt = f"Booked in other OPD\n{sheet} — {day} {dt} — {period}\nPreceptor: {pre}"
                            try:
                                cell.comment = Comment(txt, "MD↔PA conflict")
                            except Exception:
                                pass

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()
