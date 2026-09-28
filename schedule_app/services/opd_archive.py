"""Encrypted original OPD storage and retrieval. The archive format and Secrets are unchanged.

Extracted from the supplied app; this module performs no page rendering on import.
"""

from cryptography.fernet import Fernet
from cryptography.fernet import InvalidToken
from cryptography.fernet import MultiFernet
from dataclasses import dataclass
from dataclasses import field
from datetime import date as CalendarDate
from datetime import datetime
from datetime import timedelta
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from urllib.parse import quote
from zipfile import ZipFile
import base64
import hashlib
import hmac
import re
import requests
import streamlit as st
import zipfile


OPD_ARCHIVE_VERSION = 1
OPD_MAX_BYTES = 10 * 1024 * 1024
OPD_MAX_ENCRYPTED_BYTES = 15 * 1024 * 1024
OPD_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
OPD_NAME_ORDERS = (
    "Auto-detect using rotation list",
    "Student ~ Preceptor",
    "Preceptor ~ Student",
)


class OPDArchiveError(RuntimeError):
    """A safe, user-facing archive error; never include a token or workbook data."""


def _opd_secret_settings():
    try:
        return dict(st.secrets.get("opd_archive", {}))
    except Exception:
        return {}


@dataclass(frozen=True)
class OPDArchiveConfig:
    owner: str
    repo: str
    branch: str
    folder: str
    github_token: str = field(repr=False)
    encryption_key: str = field(repr=False)
    previous_encryption_keys: tuple = field(default=(), repr=False)

    def cipher(self):
        try:
            keys = (self.encryption_key,) + self.previous_encryption_keys
            return MultiFernet([Fernet(k.encode("ascii")) for k in keys])
        except (ValueError, TypeError, UnicodeError):
            raise OPDArchiveError("An archive encryption key is invalid. Use a Fernet "
                                  "key from generate_opd_secrets.py.") from None

    def signature(self):
        # Session-only identity. Includes credential changes to invalidate receipts.
        material = "\0".join((self.owner, self.repo, self.branch, self.folder,
                              self.github_token, self.encryption_key,
                              *self.previous_encryption_keys))
        return hashlib.sha256(material.encode()).hexdigest()


def get_opd_archive_config():
    values = _opd_secret_settings()
    required = ("owner", "repo", "github_token", "encryption_key")
    if any(not str(values.get(k, "")).strip() for k in required):
        raise OPDArchiveError("Archive setup is incomplete. Add owner, repo, "
                              "github_token and encryption_key under [opd_archive] "
                              "in Streamlit Secrets.")
    owner, repo = str(values["owner"]).strip(), str(values["repo"]).strip()
    if not re.fullmatch(r"[A-Za-z0-9-]+", owner) or not re.fullmatch(r"[A-Za-z0-9_.-]+", repo):
        raise OPDArchiveError("Use the GitHub account name and repository name only, not a URL.")
    branch = str(values.get("branch", "main")).strip()
    folder = str(values.get("folder", "opd_archive")).strip("/")
    if not branch or not folder or any(
        not re.fullmatch(r"[A-Za-z0-9_-]+", part) for part in folder.split("/")
    ):
        raise OPDArchiveError("Check the archive branch and folder in Streamlit Secrets.")
    previous = values.get("previous_encryption_keys", [])
    if not isinstance(previous, (tuple, list)):
        raise OPDArchiveError("previous_encryption_keys must be a TOML list, or omitted.")
    config = OPDArchiveConfig(owner, repo, branch, folder,
                              str(values["github_token"]).strip(),
                              str(values["encryption_key"]).strip(),
                              tuple(str(k).strip() for k in previous))
    config.cipher()
    return config


def validate_opd_xlsx_bytes(raw):
    """Bound memory use and reject non-XLSX inputs without modifying the bytes."""
    if not isinstance(raw, bytes) or not raw:
        raise OPDArchiveError("The OPD upload is empty.")
    if len(raw) > OPD_MAX_BYTES:
        raise OPDArchiveError("OPD uploads must be 10 MB or smaller.")
    try:
        with ZipFile(BytesIO(raw)) as zf:
            members = zf.infolist()
            if len(members) > 20000 or sum(i.file_size for i in members) > 100 * 1024 * 1024:
                raise OPDArchiveError("This workbook is too large to process safely.")
            names = set(zf.namelist())
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise OPDArchiveError("Upload an actual Excel .xlsx workbook, not an .xls file.")
            if any(i.flag_bits & 1 for i in members) or "xl/vbaProject.bin" in names:
                raise OPDArchiveError("Use a normal .xlsx file without password protection or macros.")
            if zf.testzip() is not None:
                raise OPDArchiveError("The uploaded workbook failed its ZIP integrity check.")
    except zipfile.BadZipFile:
        raise OPDArchiveError("The file could not be read as an Excel .xlsx workbook.") from None


def _opd_date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, CalendarDate):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            dt = from_excel(value, epoch)
            return dt.date() if isinstance(dt, datetime) and 1970 <= dt.year <= 2100 else None
        except (ValueError, OverflowError):
            return None
    if isinstance(value, str):
        for fmt in ("%m/%d/%Y", "%m-%d-%Y", "%Y-%m-%d", "%B %d, %Y", "%b %d, %Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def inspect_opd_rotation(raw):
    """Find the first scheduled Monday in site-tab OPDs, using calendar headers.

    The filename and names around '~' play no role in the archive identifier.
    Different date ranges across site worksheets are rejected, not guessed.
    """
    validate_opd_xlsx_bytes(raw)
    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise OPDArchiveError("The OPD workbook could not be opened. Save it as .xlsx and retry.") from None
    try:
        expected_days = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
        schedules = {}
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 1200),
                                     max_col=8, values_only=True))
            if not rows or str(rows[0][0] or "").strip().rstrip(":").casefold() != "site":
                continue
            has_sessions = any(re.match(r"^\s*(AM|PM)\b", str(row[0] or ""), re.I) for row in rows)
            if not has_sessions:
                continue
            mondays = []
            for r, row in enumerate(rows[:-1]):
                labels = [str(v or "").strip().casefold() for v in row[1:8]]
                if labels != expected_days:
                    continue
                values = [_opd_date(v, wb.epoch) for v in rows[r + 1][1:8]]
                if not all(values) or values[0].weekday() != 0 or any(
                    day != values[0] + timedelta(days=i) for i, day in enumerate(values)
                ):
                    raise OPDArchiveError(f"The date row on worksheet '{ws.title}' near row {r+2} "
                                          "must contain seven consecutive Monday-Sunday dates.")
                mondays.append(values[0])
            if not mondays:
                raise OPDArchiveError(f"No readable Monday-Sunday date headers were found on '{ws.title}'.")
            if any(d != mondays[0] + timedelta(days=7*i) for i, d in enumerate(mondays)):
                raise OPDArchiveError(f"The calendar weeks on '{ws.title}' are not consecutive.")
            schedules[ws.title] = tuple(mondays)
        if not schedules:
            raise OPDArchiveError("This is not a recognized OPD site-tab workbook. Upload the original "
                                  "OPD.xlsx, not the one-tab-per-student MS_Schedule.xlsx.")
        first = next(iter(schedules.values()))
        if any(weeks != first for weeks in schedules.values()):
            raise OPDArchiveError("The OPD site worksheets have different rotation dates. "
                                  "Correct the dates before archiving; no file was replaced.")
        return {"rotation_start": first[0], "week_mondays": first,
                "site_names": tuple(schedules), "sha256": hashlib.sha256(raw).hexdigest()}
    finally:
        wb.close()


class GitHubOPDArchive:
    """GitHub Contents API transport; only authenticated ciphertext is written.

    Reads are pinned to a commit. Updates use the existing blob SHA. A competing
    write causes an explicit retry prompt rather than a silent overwrite/rebase.
    """
    def __init__(self, config, transport=None):
        self.config = config
        self.transport = transport or requests
        self.base = f"https://api.github.com/repos/{config.owner}/{config.repo}"
        self.cipher = config.cipher()

    def _request(self, method, route, *, params=None, body=None, raw=False, missing_ok=False):
        headers = {"Authorization": f"Bearer {self.config.github_token}",
                   "Accept": ("application/vnd.github.raw+json" if raw else
                              "application/vnd.github.object+json" if method == "GET" else "application/vnd.github+json"),
                   "X-GitHub-Api-Version": "2026-03-10",
                   "User-Agent": "OPD-Encrypted-Archive/1.0"}
        try:
            response = self.transport.request(method, self.base + route, headers=headers,
                                              params=params, json=body, timeout=(10, 45),
                                              allow_redirects=False, stream=raw)
        except requests.RequestException:
            raise OPDArchiveError("GitHub could not be reached. The save is not confirmed. "
                                  "Retry when the connection is available.") from None
        status = response.status_code
        if status == 404 and missing_ok:
            response.close()
            return None
        if status not in (200, 201):
            response.close()
            messages = {
                401: "GitHub authentication failed. Check or renew the access token in Streamlit Secrets.",
                403: "GitHub denied access or rate-limited this request. Check token permissions, organization approval and repository rules, then retry later.",
                404: "GitHub repository, branch or archive file was not found. Check owner, repo and branch.",
                409: "Another save changed the repository during this operation. Retry explicitly; the app will re-read the current version first.",
                422: "GitHub rejected the write. Check branch rules and token permissions, then retry.",
                429: "GitHub rate-limited this request. Wait before retrying.",
            }
            raise OPDArchiveError(messages.get(status, f"GitHub request failed (HTTP {status}); the save is not confirmed."))
        if raw:
            try:
                chunks, size = [], 0
                for chunk in response.iter_content(chunk_size=65536):
                    size += len(chunk)
                    if size > OPD_MAX_ENCRYPTED_BYTES:
                        raise OPDArchiveError("The encrypted archive file exceeds this app's size limit.")
                    chunks.append(chunk)
                return b"".join(chunks)
            except requests.RequestException:
                raise OPDArchiveError("The encrypted file download was interrupted. Retry loading it.") from None
            finally:
                response.close()
        try:
            return response.json()
        except (ValueError, requests.RequestException):
            raise OPDArchiveError("GitHub returned an unexpected response. No save is confirmed.") from None
        finally:
            response.close()

    def _head(self):
        data = self._request("GET", "/branches/" + quote(self.config.branch, safe=""))
        try:
            return data["commit"]["sha"]
        except (KeyError, TypeError):
            raise OPDArchiveError("GitHub did not return a valid branch. Initialize the archive repository with a README.") from None

    def path_for(self, rotation_start):
        if not isinstance(rotation_start, CalendarDate) or rotation_start.weekday() != 0:
            raise OPDArchiveError("The archive identifier must be the rotation's first Monday.")
        return f"{self.config.folder}/OPD_{rotation_start.isoformat()}.xlsx.enc"

    def _read_at(self, path, commit):
        route = "/contents/" + quote(path, safe="/")
        metadata = self._request("GET", route, params={"ref": commit}, missing_ok=True)
        if metadata is None:
            return None
        if not isinstance(metadata, dict) or metadata.get("type") != "file" or metadata.get("submodule_git_url"):
            raise OPDArchiveError("The archive path is not a regular file. No replacement was made.")
        if int(metadata.get("size", 0)) > OPD_MAX_ENCRYPTED_BYTES:
            raise OPDArchiveError("The encrypted archive file is too large for this app.")
        if metadata.get("encoding") == "base64" and metadata.get("content"):
            try:
                token = base64.b64decode("".join(metadata["content"].split()), validate=True)
            except (ValueError, TypeError):
                raise OPDArchiveError("GitHub returned invalid encoded file contents.") from None
        else:
            token = self._request("GET", route, params={"ref": commit}, raw=True)
        if len(token) > OPD_MAX_ENCRYPTED_BYTES:
            raise OPDArchiveError("The encrypted archive file is too large for this app.")
        blob_sha = hashlib.sha1(b"blob " + str(len(token)).encode() + b"\0" + token).hexdigest()
        if blob_sha != metadata.get("sha"):
            raise OPDArchiveError("The archive download did not match its GitHub file identifier. Retry loading it.")
        try:
            plaintext = self.cipher.decrypt(token)  # no TTL: historical files remain reloadable
        except (InvalidToken, ValueError):
            raise OPDArchiveError("This OPD cannot be decrypted with the configured key(s), or its encrypted "
                                  "contents were altered. The existing file will NOT be overwritten. "
                                  "Restore the correct encryption key in Streamlit Secrets.") from None
        details = inspect_opd_rotation(plaintext)
        if path != self.path_for(details["rotation_start"]):
            raise OPDArchiveError("The decrypted OPD's rotation does not match its archive filename.")
        return {"raw": plaintext, "encrypted": token, "details": details,
                "sha": metadata["sha"], "path": path, "commit": commit}

    def load(self, rotation_start, commit=None):
        found = self._read_at(self.path_for(rotation_start), commit or self._head())
        if found is None:
            raise OPDArchiveError("No archived OPD was found for that rotation. Refresh the archive list.")
        return found

    def save(self, raw):
        details = inspect_opd_rotation(raw)
        path = self.path_for(details["rotation_start"])
        current = self._read_at(path, self._head())
        if current is not None and hmac.compare_digest(current["raw"], raw):
            return {"action": "unchanged", "path": path, "sha": current["sha"],
                    "sha256": details["sha256"], "rotation_start": details["rotation_start"]}
        token = self.cipher.encrypt(raw)
        body = {"message": "Update encrypted OPD archive", "branch": self.config.branch,
                "content": base64.b64encode(token).decode("ascii")}
        if current is not None:
            body["sha"] = current["sha"]
        result = self._request("PUT", "/contents/" + quote(path, safe="/"), body=body)
        try:
            commit = result["commit"]["sha"]
        except (KeyError, TypeError):
            raise OPDArchiveError("GitHub did not return a save confirmation. Retry to verify the current archive.") from None
        verified = self._read_at(path, commit)
        if verified is None or not hmac.compare_digest(verified["raw"], raw):
            raise OPDArchiveError("The uploaded OPD could not be verified after saving. Retry before creating schedules.")
        return {"action": "replaced" if current else "created", "path": path,
                "sha": verified["sha"], "sha256": details["sha256"],
                "rotation_start": details["rotation_start"]}

    def list_rotations(self, commit=None):
        commit = commit or self._head()
        route = "/contents/" + quote(self.config.folder, safe="/")
        data = self._request("GET", route, params={"ref": commit}, missing_ok=True)
        if data is None:
            return []
        entries = data.get("entries") if isinstance(data, dict) else data
        if not isinstance(entries, list):
            raise OPDArchiveError("The configured archive folder is not a directory.")
        if len(entries) >= 1000:
            raise OPDArchiveError("The GitHub directory listing limit was reached. The app cannot safely "
                                  "show a complete archive; organize the archive before continuing.")
        rotations = []
        for item in entries:
            match = re.fullmatch(r"OPD_(\d{4}-\d{2}-\d{2})\.xlsx\.enc", str(item.get("name", "")))
            if item.get("type") != "file" or not match:
                continue
            try:
                day = CalendarDate.fromisoformat(match.group(1))
            except ValueError:
                continue
            if day.weekday() == 0:
                rotations.append(day)
        return sorted(set(rotations), reverse=True)
