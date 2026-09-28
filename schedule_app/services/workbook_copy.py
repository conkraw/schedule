"""Copy individual student worksheets with the existing formatting behavior.

Extracted from the supplied app; this module performs no page rendering on import.
"""

from copy import copy
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string


def copy_sheet_to_new_wb(src_ws):
    """Return a BytesIO of a new .xlsx containing src_ws with formatting."""
    from openpyxl import Workbook
    from io import BytesIO

    wb_new = Workbook()
    ws_new = wb_new.active
    ws_new.title = src_ws.title[:31]

    # Column widths & visibility
    for col_letter, dim in src_ws.column_dimensions.items():
        if dim.width is not None:
            ws_new.column_dimensions[col_letter].width = dim.width
        ws_new.column_dimensions[col_letter].hidden = dim.hidden

    # Row heights & visibility
    for idx, dim in src_ws.row_dimensions.items():
        if dim.height is not None:
            ws_new.row_dimensions[idx].height = dim.height
        ws_new.row_dimensions[idx].hidden = dim.hidden

    # Sheet settings (best-effort)
    try:
        ws_new.sheet_format.defaultColWidth = src_ws.sheet_format.defaultColWidth
        ws_new.sheet_format.defaultRowHeight = src_ws.sheet_format.defaultRowHeight
    except Exception:
        pass
    ws_new.freeze_panes = src_ws.freeze_panes
    try:
        ws_new.page_setup.orientation = src_ws.page_setup.orientation
        ws_new.page_setup.fitToWidth = src_ws.page_setup.fitToWidth
        ws_new.page_setup.fitToHeight = src_ws.page_setup.fitToHeight
        ws_new.page_margins = copy(src_ws.page_margins)
        ws_new.print_options.horizontalCentered = src_ws.print_options.horizontalCentered
        ws_new.print_options.verticalCentered = src_ws.print_options.verticalCentered
        ws_new.print_area = src_ws.print_area
    except Exception:
        pass

    try:
        # Set zoom to 70%
        ws_new.sheet_view.zoomScale = 70

        # Set column widths
        ws_new.column_dimensions["A"].width = 20
        ws_new.column_dimensions["B"].width = 30
        for col in ["C", "D", "E", "F", "G"]:
            ws_new.column_dimensions[col].width = 40
        ws_new.column_dimensions["H"].width = 155
    except Exception:
        pass

    # Copy cells: values + (copied) styles
    for row in src_ws.iter_rows():
        for cell in row:
            # Skip non-master cells from merged ranges
            if isinstance(cell, MergedCell):
                continue

            # Get a reliable numeric column index
            col_idx = getattr(cell, "col_idx", None)
            if col_idx is None:
                col = cell.column  # may be int or letter depending on version
                col_idx = col if isinstance(col, int) else column_index_from_string(col)

            # Create target cell with value (formula preserved if present)
            tgt = ws_new.cell(row=cell.row, column=col_idx, value=cell.value)

            # Copy style safely
            if getattr(cell, "has_style", False):
                try:
                    if cell.font:
                        tgt.font = copy(cell.font)
                    if cell.fill:
                        tgt.fill = copy(cell.fill)
                    if cell.border:
                        tgt.border = copy(cell.border)
                    if cell.alignment:
                        tgt.alignment = copy(cell.alignment)
                    if cell.protection:
                        tgt.protection = copy(cell.protection)
                    tgt.number_format = cell.number_format
                except Exception:
                    pass

    # Copy merged cell ranges (after values)
    for merged in list(src_ws.merged_cells.ranges):
        try:
            ws_new.merge_cells(str(merged))
        except Exception:
            pass

    # Copy data validations (best-effort)
    try:
        if src_ws.data_validations and src_ws.data_validations.dataValidation:
            from openpyxl.worksheet.datavalidation import DataValidation

            for dv in src_ws.data_validations.dataValidation:
                dv_new = DataValidation(
                    type=dv.type,
                    formula1=dv.formula1,
                    formula2=dv.formula2,
                    allow_blank=dv.allow_blank,
                    operator=dv.operator,
                    showDropDown=dv.showDropDown,
                    showErrorMessage=dv.showErrorMessage,
                    errorTitle=dv.errorTitle,
                    error=dv.error,
                    promptTitle=dv.promptTitle,
                    prompt=dv.prompt,
                )
                for sqref in getattr(dv, "sqref", []):
                    dv_new.add(sqref)
                ws_new.add_data_validation(dv_new)
    except Exception:
        pass

    # Filters
    try:
        ws_new.auto_filter.ref = getattr(src_ws.auto_filter, "ref", None)
    except Exception:
        pass

    # Save to buffer
    buf = BytesIO()
    wb_new.save(buf)
    buf.seek(0)
    return buf
