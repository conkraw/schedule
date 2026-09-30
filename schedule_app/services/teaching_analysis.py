"""Read-only archive analysis: student-shifts, academic years, deduplication and work types.

Raw no_of_shifts and raw educational_hours in a scan remain legacy per-student
assignment aggregates for source/continuity integrity checks. They are never
published as preceptor time. Public annual and work-type summaries derive
educational_hours from the unique clinical shifts with students in learner_reach.
This module performs no page rendering on import.
"""

from collections import Counter
from collections import defaultdict
from datetime import date as CalendarDate
from datetime import datetime
from datetime import timedelta
from datetime import timezone as _teaching_timezone
from io import BytesIO
from openpyxl import load_workbook
from schedule_app.services.opd_archive import OPDArchiveError
from schedule_app.services.opd_archive import _opd_date
from schedule_app.settings import TEACHING_EMPTY_STUDENT_LABELS
from schedule_app.settings import TEACHING_HOURS_PER_STUDENT_SHIFT
from schedule_app.settings import TEACHING_MONTH_NAMES
from schedule_app.settings import TEACHING_NAME_ORDERS
from schedule_app.settings import TEACHING_OPD_NAME_ORDER_OVERRIDES
from schedule_app.settings import TEACHING_PRECEPTOR_NAME_MAP
from schedule_app.settings import TEACHING_REPORT_VERSION
from schedule_app.settings import TEACHING_WORK_TYPE_MAP
from schedule_app.settings import TEACHING_WORK_TYPE_ORDER
from schedule_app.settings import TEACHING_WORK_TYPE_REVIEW
from zoneinfo import ZoneInfo as _TeachingZoneInfo
import hashlib
import hmac
import json
import re
import secrets as _teaching_secrets
from schedule_app.services.reporting_periods import (
    DATE_RANGE_SCHEMA_VERSION, ReportingPeriod, teaching_report_label,
)


from schedule_app.services.learner_reach import (
    ClinicalShiftAccumulator, enrich_teaching_rows, filter_reach_dates,
    nonclinical_provider, require_learner_reach_data,
)

from schedule_app.services.student_continuity import (
    build_student_continuity_data, require_student_continuity_data,
    filter_student_continuity_dates,
)


def teaching_site_key(site):
    return re.sub(r"[^A-Z0-9]+", "_", str(site or "").upper()).strip("_")


def teaching_work_type(site):
    """Classify the assignment by its OPD site, not the preceptor's name."""
    key = teaching_site_key(site)
    if not key:
        raise OPDArchiveError("An assigned shift has no OPD site; its work type cannot be determined.")
    if key in TEACHING_WORK_TYPE_MAP:
        label = str(TEACHING_WORK_TYPE_MAP[key]).strip()
        if not label or label == TEACHING_WORK_TYPE_REVIEW:
            raise OPDArchiveError("Check TEACHING_WORK_TYPE_MAP: use a nonempty work-type label, not the reserved review label.")
        return label
    # Preserve unmapped services, rather than silently assuming inpatient or outpatient.
    return key.replace("_", " ")


def teaching_work_type_sort(label):
    if label in TEACHING_WORK_TYPE_ORDER:
        return (0, TEACHING_WORK_TYPE_ORDER.index(label), "")
    if label == TEACHING_WORK_TYPE_REVIEW:
        return (2, 0, "")
    return (1, 0, str(label).casefold())


def teaching_require_work_type_data(scan):
    """Reject cached scans without site detail; totals alone cannot reconstruct it."""
    if scan.get("version") != TEACHING_REPORT_VERSION or "monthly_by_work_type" not in scan:
        raise OPDArchiveError("This teaching scan predates work-type reporting. Click Load / refresh archived OPDs, then generate the reports again.")
    expected, actual = Counter(), Counter()
    for row in scan["monthly"]:
        expected[(row["preceptor_name"], row["academic_start_year"], row["month"])] += row["no_of_shifts"]
    for row in scan["monthly_by_work_type"]:
        if not row.get("work_type") or int(row["no_of_shifts"]) <= 0:
            raise OPDArchiveError("The work-type breakdown is incomplete. Refresh the archived OPDs; no partial report was generated.")
        actual[(row["preceptor_name"], row["academic_start_year"], row["month"])] += row["no_of_shifts"]
    if dict(expected) != dict(actual):
        raise OPDArchiveError("Work-type subtotals do not match the overall teaching totals. Refresh the archived OPDs before generating a report.")


def teaching_work_type_rows(scan, selected_years):
    """One row per preceptor, academic year and work type; no student identifiers."""
    teaching_require_work_type_data(scan)
    years = {int(year) for year in selected_years}
    grouped = defaultdict(list)
    for item in scan["monthly_by_work_type"]:
        if item["academic_start_year"] in years:
            grouped[(item["preceptor_name"], item["academic_start_year"], item["work_type"])].append(item)
    rows = []
    for (name, year, work_type), items in sorted(
        grouped.items(), key=lambda pair: (teaching_name_key(pair[0][0]), pair[0][1], teaching_work_type_sort(pair[0][2]))
    ):
        months = sorted({item["month"] for item in items})
        total = sum(item["no_of_shifts"] for item in items)
        sites = sorted({site for item in items for site in item["source_sites"]})
        rows.append({
            "preceptor_name": name, "academic_year": teaching_report_label(scan, year),
            "work_type": work_type, "no_of_shifts": total,
            "months_worked": "; ".join(teaching_month_label(CalendarDate.fromisoformat(month)) for month in months),
            "educational_hours": total * TEACHING_HOURS_PER_STUDENT_SHIFT,
            "source_sites": "; ".join(sites),
        })
    require_learner_reach_data(scan)
    return enrich_teaching_rows(scan, selected_years, rows, by_work_type=True)


def teaching_name_key(value):
    cleaned = re.sub(r"\s+", " ", str(value or "").strip())
    return re.sub(r"\s*,\s*", ", ", cleaned).casefold()


def teaching_display_name(value):
    return re.sub(r"\s*,\s*", ", ", re.sub(r"\s+", " ", str(value or "").strip()))


def teaching_local_today():
    # Streamlit Cloud may run on UTC. Keep the academic-year boundary local.
    try:
        return datetime.now(_TeachingZoneInfo("America/New_York")).date()
    except Exception:
        return datetime.now(_teaching_timezone.utc).date()


def teaching_academic_start(day):
    return day.year if day.month >= 7 else day.year - 1


def teaching_academic_label(start_year):
    return f"{start_year % 100:02d}-{(start_year + 1) % 100:02d}"


def teaching_month_label(month):
    return f"{TEACHING_MONTH_NAMES[month.month - 1]} {month.year}"


def teaching_label_needs_review(name, site_names=()):
    key = re.sub(r"[^A-Z0-9]+", "_", str(name).upper()).strip("_")
    site_keys = {re.sub(r"[^A-Z0-9]+", "_", str(s).upper()).strip("_") for s in site_names}
    return (
        key in site_keys or key in {"HAMPDEN_NURSERY", "PSHCH_NURSERY", "SJR_HOSPITALIST", "WARD_A"}
        or bool(re.fullmatch(r"(?:SJR|AAC|HOPE_DRIVE|NYES|ETOWN|LANCASTER|HAMPDEN_NURSERY|PSHCH_NURSERY)_?\d+", key))
        or key in {"TBD", "UNKNOWN", "PRECEPTOR", "PROVIDER", "TO_BE_ASSIGNED"}
        or bool(re.search(r"[;&|]|\s/\s|\s+and\s+", str(name), re.I))
    )


def teaching_split_assignment(value, order):
    """Return (provider, [students]) groups without splitting 'Last, First'.

    Names are used only temporarily to count separate students and recognize
    exact duplicates. They never enter the exported CSV, Word files or notes.
    Multiple students: separate OPD rows, or ; / newline / | / ' & ' / ' and '
    on the student side. Repeated full Provider ~ Student pairs may be separated
    with a newline, semicolon or |. A bare repeated ~ is ambiguous and rejected.
    """
    if order not in TEACHING_NAME_ORDERS:
        raise OPDArchiveError("Choose Preceptor ~ Student or Student ~ Preceptor for the teaching summary.")
    if not isinstance(value, str) or "~" not in value:
        return []
    text = value.strip()
    if text.count("~") == 1:
        segments = [text]
    else:
        segments = [part.strip() for part in re.split(r"[\r\n;|]+", text) if part.strip()]
        if any(part.count("~") != 1 for part in segments):
            raise OPDArchiveError("An assignment contains an ambiguous repeated '~'. Use separate OPD rows "
                                  "or a semicolon-separated list of students on the student side.")
    parsed = []
    for segment in segments:
        left, right = (part.strip() for part in segment.split("~", 1))
        provider, student_text = (left, right) if order == "Preceptor ~ Student" else (right, left)
        students = [teaching_display_name(part) for part in
                    re.split(r"[\r\n;|]+|\s+&\s+|\s+and\s+", student_text, flags=re.I)]
        students = [name for name in students if teaching_name_key(name) not in (set(TEACHING_EMPTY_STUDENT_LABELS) | {"-", "--", "—", "null"})]
        parsed.append((teaching_display_name(provider), students))
    return parsed


def teaching_extract_assignments(raw, details, order, *, clinical_sessions=None, ignored_cells=None):
    """Read AM/PM cells under each *actual* date, including hidden OPD rows.

    Returns temporary records containing student names. The caller aggregates
    and discards these records before storing anything in Streamlit session_state.
    """
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    records, missing_provider_cells = [], []
    # Optional output lists preserve the existing extractor's two-value return.
    # Clinical records never contain learner identifiers.
    if clinical_sessions is None:
        clinical_sessions = []
    if ignored_cells is None:
        ignored_cells = []
    expected_days = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
    try:
        for sheet_name in details["site_names"]:
            ws = wb[sheet_name]
            if (ws.max_row or 0) > 1200:
                raise OPDArchiveError(f"Worksheet '{sheet_name}' exceeds the supported 1,200-row OPD layout; "
                                      "no partial teaching report was generated.")
            rows = list(ws.iter_rows(max_col=8))
            current_dates = None
            for row_index, row in enumerate(rows):
                values = [cell.value for cell in row]
                if [str(v or "").strip().casefold() for v in values[1:8]] == expected_days:
                    if row_index + 1 >= len(rows):
                        raise OPDArchiveError(f"Missing dates on worksheet '{sheet_name}'.")
                    current_dates = [_opd_date(cell.value, wb.epoch) for cell in rows[row_index + 1][1:8]]
                    if not all(current_dates) or any(
                        day != current_dates[0] + timedelta(days=i) for i, day in enumerate(current_dates)
                    ):
                        raise OPDArchiveError(f"Invalid date row on worksheet '{sheet_name}'.")
                    continue
                match = re.match(r"^\s*(AM|PM)\b", str(values[0] or ""), re.I)
                if not match:
                    continue
                for col_index, cell in enumerate(row[1:8]):
                    if cell.data_type == "f":
                        raise OPDArchiveError(f"{sheet_name}!{cell.coordinate} contains a formula in a session cell. "
                                              "Use assignment values rather than formulas for this summary.")
                    if cell.value is not None and str(cell.value).strip() and "~" not in str(cell.value):
                        ignored_cells.append({"cell": f"{sheet_name}!{cell.coordinate}", "reason": "No '~' marker"})
                    try:
                        groups = teaching_split_assignment(cell.value, order)
                    except OPDArchiveError as exc:
                        raise OPDArchiveError(f"{sheet_name}!{cell.coordinate}: {exc}") from None
                    for provider, students in groups:
                        if teaching_name_key(provider) in TEACHING_EMPTY_STUDENT_LABELS:
                            if students:
                                missing_provider_cells.append(f"{sheet_name}!{cell.coordinate}")
                            elif provider:
                                ignored_cells.append({"cell": f"{sheet_name}!{cell.coordinate}", "reason": "Empty/nonclinical provider label"})
                            continue
                        if nonclinical_provider(provider):
                            ignored_cells.append({"cell": f"{sheet_name}!{cell.coordinate}", "reason": "Nonclinical provider label"})
                            if students:
                                missing_provider_cells.append(f"{sheet_name}!{cell.coordinate}")
                            continue
                        if current_dates is None:
                            raise OPDArchiveError(f"{sheet_name}!{cell.coordinate} has a provider listing before a date header.")
                        clinical_sessions.append({
                            "preceptor_name": provider, "has_student": bool(students),
                            "day": current_dates[col_index], "shift": match.group(1).upper(),
                            "site": sheet_name, "cell": cell.coordinate,
                        })
                        for student in students:
                            records.append({
                                "preceptor_name": provider, "student": student,
                                "day": current_dates[col_index], "shift": match.group(1).upper(),
                                "site": sheet_name, "cell": cell.coordinate,
                            })
        return records, missing_provider_cells
    finally:
        wb.close()


def teaching_scan_archives(client, default_order=TEACHING_NAME_ORDERS[0], progress=None):
    """Read/decrypt each current OPD at one repository snapshot, retaining work type.

    Student names are temporary. Run-local keyed digests deduplicate assignments
    and group distinct teaching dates per learner. Only unlinked date groups are
    retained in session memory; digests/names never enter returned scans/reports.
    Work-type subtotals always reconcile to the existing overall counting rules.
    """
    if default_order not in TEACHING_NAME_ORDERS:
        raise OPDArchiveError("Select a supported OPD name order.")
    commit = client._head()
    rotations = client.list_rotations(commit=commit)
    counts = Counter()
    names, name_variants, review_labels, manifests, warnings = {}, defaultdict(set), set(), [], []
    # One anonymous item per provider/student/date/AM-or-PM assignment.
    seen_assignments = {}
    clinical = ClinicalShiftAccumulator()
    salt = _teaching_secrets.token_bytes(32)
    aliases = {teaching_name_key(k): teaching_display_name(v)
               for k, v in TEACHING_PRECEPTOR_NAME_MAP.items() if str(k).strip() and str(v).strip()}
    duplicates = future_assignments = 0
    today = teaching_local_today()
    observed_site_groups = {}
    for number, rotation in enumerate(rotations, start=1):
        order = TEACHING_OPD_NAME_ORDER_OVERRIDES.get(rotation.isoformat(), default_order)
        if order not in TEACHING_NAME_ORDERS:
            raise OPDArchiveError(f"Invalid teaching-summary name-order override for rotation {rotation.isoformat()}.")
        try:
            loaded = client.load(rotation, commit=commit)
            clinical_sessions, ignored_cells = [], []
            records, missing_cells = teaching_extract_assignments(
                loaded["raw"], loaded["details"], order,
                clinical_sessions=clinical_sessions, ignored_cells=ignored_cells,
            )
        except OPDArchiveError as exc:
            raise OPDArchiveError(f"Rotation {rotation.isoformat()}: {exc} No ZIP was generated from partial data.") from None
        counted = removed = 0
        for item in records:
            original_name = item["preceptor_name"]
            name = aliases.get(teaching_name_key(original_name), original_name)
            key = teaching_name_key(name)
            work_type = teaching_work_type(item["site"])
            site = teaching_site_key(item["site"])
            observed_site_groups[site] = work_type
            identity = json.dumps([key, teaching_name_key(item["student"]),
                                   item["day"].isoformat(), item["shift"]], ensure_ascii=False)
            digest = hmac.new(salt, identity.encode("utf-8"), hashlib.sha256).digest()
            # A duplicate within Academic Pediatrics is still counted once even
            # if it is copied onto both HOPE_DRIVE and NYES worksheets. Distinct
            # students in the same shift have different identities and count twice.
            origin = (rotation.isoformat(), site)
            if digest in seen_assignments:
                seen_assignments[digest]["origin_counts"][origin] += 1
                duplicates += 1
                removed += 1
                seen_assignments[digest]["work_types"].add(work_type)
                seen_assignments[digest]["sites"].add(site)
                continue
            month = item["day"].replace(day=1)
            # A separate, domain-scoped identifier links this learner's retained
            # dates across rotations/work types. It is discarded before return.
            student_identity = json.dumps(
                ["student-continuity", key, teaching_name_key(item["student"])],
                ensure_ascii=False,
            )
            student_key = hmac.new(salt, student_identity.encode("utf-8"), hashlib.sha256).digest()
            seen_assignments[digest] = {
                "student_key": student_key,
                "provider_key": key, "month": month, "day": item["day"], "shift": item["shift"],
                "work_types": {work_type}, "sites": {site},
                "origin_counts": Counter({origin: 1}),
            }
            names.setdefault(key, name)
            name_variants[key].add(original_name)
            counts[(key, month)] += 1
            counted += 1
            future_assignments += int(item["day"] > today)
            if teaching_label_needs_review(name, loaded["details"]["site_names"]):
                review_labels.add(key)
        for item in clinical_sessions:
            original_name = item["preceptor_name"]
            name = aliases.get(teaching_name_key(original_name), original_name)
            key = teaching_name_key(name)
            names.setdefault(key, name)
            name_variants[key].add(original_name)
            site = teaching_site_key(item["site"])
            work_type = teaching_work_type(item["site"])
            observed_site_groups[site] = work_type
            clinical.add(key, item, work_type, site, source={
                "rotation_start": rotation.isoformat(), "archive_path": loaded["path"],
                "archive_file": loaded["path"].rsplit("/", 1)[-1],
                "github_blob_sha": loaded["sha"], "worksheet": item["site"], "cell": item["cell"],
            })
            if teaching_label_needs_review(name, loaded["details"]["site_names"]):
                review_labels.add(key)
        for reason in sorted({row["reason"] for row in ignored_cells}):
            warnings.append({"rotation_start": rotation.isoformat(), "issue": reason + "; excluded",
                             "details": ", ".join(sorted({row["cell"] for row in ignored_cells if row["reason"] == reason}))})
        if missing_cells:
            warnings.append({"rotation_start": rotation.isoformat(), "issue": "Missing provider; not attributed",
                             "details": ", ".join(sorted(set(missing_cells)))})
        manifests.append({
            "rotation_start": rotation.isoformat(),
            "last_scheduled_date": (loaded["details"]["week_mondays"][-1] + timedelta(days=6)).isoformat(),
            "archive_file": loaded["path"].rsplit("/", 1)[-1],
            "github_blob_sha": loaded["sha"], "name_order": order,
            "assigned_student_shifts_read": len(records),
            "assigned_student_shifts_counted": counted,
            "duplicate_student_shifts_removed": removed,
            "missing_provider_cells": len(set(missing_cells)),
            "clinical_provider_listings_read": len(clinical_sessions),
            "ignored_session_cells": len({row["cell"] for row in ignored_cells}),
        })
        del records, loaded, clinical_sessions, ignored_cells
        if progress:
            progress(number, len(rotations))
    # Apply the approved exception to BOTH denominators and student assignments.
    # Do this after scanning all files: overlapping rotations may contain the
    # outpatient listing in a different file from the nursery listing.
    clinical.apply_outpatient_priority(names)
    counts = Counter()
    duplicates = future_assignments = 0
    source_by_rotation = {row["rotation_start"]: row for row in manifests}
    for row in manifests:
        row["assigned_student_shifts_counted"] = 0
        row["duplicate_student_shifts_removed"] = 0
        row["nursery_student_assignment_listings_excluded"] = 0
    effective_assignments = {}
    removed_assignments = 0
    for digest, item in seen_assignments.items():
        identity = (item["provider_key"], item["day"], item["shift"])
        excluded = clinical.excluded_sites_by_shift.get(identity, set())
        remaining_origins = []
        for (rotation_key, site), count in item["origin_counts"].items():
            if site in excluded:
                source_by_rotation[rotation_key]["nursery_student_assignment_listings_excluded"] += count
            else:
                remaining_origins.append((rotation_key, site, count))
        if not remaining_origins:
            removed_assignments += 1
            continue  # A nursery-only student is NOT credited to clinic.
        item["sites"].difference_update(excluded)
        item["work_types"] = {observed_site_groups[site] for site in item["sites"]}
        # Reassign first-source credit only if that source was excluded. Each
        # retained student-shift still counts once across all duplicate listings.
        for index, (rotation_key, site, count) in enumerate(remaining_origins):
            credited = int(index == 0)
            source_by_rotation[rotation_key]["assigned_student_shifts_counted"] += credited
            source_by_rotation[rotation_key]["duplicate_student_shifts_removed"] += count - credited
            duplicates += count - credited
        counts[(item["provider_key"], item["month"])] += 1
        future_assignments += int(item["day"] > today)
        del item["origin_counts"]
        effective_assignments[digest] = item
    seen_assignments = effective_assignments
    monthly = [{"preceptor_name": names[key], "academic_year": teaching_academic_label(teaching_academic_start(month)),
                "academic_start_year": teaching_academic_start(month), "month": month.isoformat(),
                "no_of_shifts": int(count), "educational_hours": int(count * TEACHING_HOURS_PER_STUDENT_SHIFT)}
               for (key, month), count in sorted(counts.items(), key=lambda item: (item[0][0], item[0][1]))]

    type_counts, conflict_counts = Counter(), Counter()
    type_sites = defaultdict(set)
    # Day-level provider aggregates allow exact custom boundaries, including
    # mid-month cutoffs. No student names or deduplication digests are retained.
    daily_counts, daily_sites = Counter(), defaultdict(set)
    for item in seen_assignments.values():
        key, month = item["provider_key"], item["month"]
        work_type = next(iter(item["work_types"])) if len(item["work_types"]) == 1 else TEACHING_WORK_TYPE_REVIEW
        type_counts[(key, month, work_type)] += 1
        type_sites[(key, month, work_type)].update(item["sites"])
        day_key = (key, item["day"], work_type)
        daily_counts[day_key] += 1
        daily_sites[day_key].update(item["sites"])
        if len(item["work_types"]) > 1:
            # Do not silently credit the same assignment to the first site read.
            # It stays in the overall total once, in an explicit review category.
            conflict_counts[(key, item["day"], item["shift"],
                             "; ".join(sorted(item["work_types"], key=teaching_work_type_sort)),
                             "; ".join(sorted(item["sites"])))] += 1
    monthly_by_type = [{
        "preceptor_name": names[key], "academic_year": teaching_academic_label(teaching_academic_start(month)),
        "academic_start_year": teaching_academic_start(month), "month": month.isoformat(),
        "work_type": work_type, "no_of_shifts": int(count),
        "educational_hours": int(count * TEACHING_HOURS_PER_STUDENT_SHIFT),
        "source_sites": sorted(type_sites[(key, month, work_type)]),
    } for (key, month, work_type), count in sorted(
        type_counts.items(), key=lambda pair: (pair[0][0], pair[0][1], teaching_work_type_sort(pair[0][2]))) ]
    conflicts = [{
        "preceptor_name": names[key], "date": day.isoformat(), "shift": shift,
        "academic_year": teaching_academic_label(teaching_academic_start(day)),
        "conflicting_work_types": types, "source_sites": sites,
        "no_of_student_shifts": count,
    } for (key, day, shift, types, sites), count in sorted(conflict_counts.items())]
    daily_by_type = [{
        "preceptor_name": names[key], "date": day.isoformat(), "work_type": work_type,
        "no_of_shifts": int(count), "educational_hours": int(count * TEACHING_HOURS_PER_STUDENT_SHIFT),
        "source_sites": sorted(daily_sites[(key, day, work_type)]),
    } for (key, day, work_type), count in sorted(
        daily_counts.items(), key=lambda pair: (pair[0][0], pair[0][1], teaching_work_type_sort(pair[0][2]))) ]
    result = {
        "student_shifts_removed_by_outpatient_priority": removed_assignments,
        "date_range_version": DATE_RANGE_SCHEMA_VERSION,
        "daily_by_work_type": daily_by_type,
        "version": TEACHING_REPORT_VERSION, "commit": commit,
        "generated_at": datetime.now(_teaching_timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "default_name_order": default_order, "monthly": monthly,
        "monthly_by_work_type": monthly_by_type, "work_type_conflicts": conflicts,
        "site_work_type_mapping": dict(sorted(observed_site_groups.items())),
        "sources": manifests, "warnings": warnings, "duplicate_assignments_removed": duplicates,
        "future_assignments_in_archive": future_assignments,
        "unresolved_preceptor_labels": sorted((names[k] for k in review_labels), key=teaching_name_key),
        "name_variants": {names[k]: sorted(v) for k, v in name_variants.items() if len(v) > 1},
    }
    # Excluded nursery learners must not contribute to unique-student counts.
    result.update(build_student_continuity_data(seen_assignments.values(), names))
    result.update(clinical.finish(names))
    require_student_continuity_data(result)
    require_learner_reach_data(result)
    teaching_require_work_type_data(result)
    return result


def teaching_annual_rows(scan, selected_years):
    selected = {int(year) for year in selected_years}
    grouped = defaultdict(list)
    for item in scan["monthly"]:
        if item["academic_start_year"] in selected:
            grouped[(item["preceptor_name"], item["academic_start_year"])].append(item)
    rows = []
    for (name, start_year), items in sorted(grouped.items(), key=lambda item: (teaching_name_key(item[0][0]), item[0][1])):
        ordered = sorted(items, key=lambda item: item["month"])
        total = sum(item["no_of_shifts"] for item in ordered)
        rows.append({"preceptor_name": name, "academic_year": teaching_report_label(scan, start_year),
                     "no_of_shifts": total,
                     "months_worked": "; ".join(teaching_month_label(CalendarDate.fromisoformat(item["month"])) for item in ordered),
                     "educational_hours": total * TEACHING_HOURS_PER_STUDENT_SHIFT})
    require_learner_reach_data(scan)
    return enrich_teaching_rows(scan, selected_years, rows)


def teaching_brief_months(month_values):
    """Compact month labels without implying teaching occurred in missing months."""
    months = sorted({CalendarDate.fromisoformat(value).replace(day=1) for value in month_values})
    if not months:
        return "Not recorded"
    groups, first, last = [], months[0], months[0]
    for month in months[1:]:
        if month.year == last.year and month.month == last.month + 1:
            last = month
        else:
            groups.append((first, last))
            first = last = month
    groups.append((first, last))
    labels = []
    for first, last in groups:
        label = TEACHING_MONTH_NAMES[first.month - 1][:3]
        if first != last:
            label += "-" + TEACHING_MONTH_NAMES[last.month - 1][:3]
        labels.append(f"{label} {first.year}")
    return "; ".join(labels)


def teaching_require_date_range_data(scan):
    """Monthly totals cannot support a mid-month cutoff; require daily aggregates."""
    teaching_require_work_type_data(scan)
    if scan.get("date_range_version") != DATE_RANGE_SCHEMA_VERSION or not isinstance(scan.get("daily_by_work_type"), list):
        raise OPDArchiveError("This saved scan does not contain exact teaching dates. Click Load / refresh archived OPDs once, then choose your reporting dates.")
    expected, actual = Counter(), Counter()
    for row in scan["monthly_by_work_type"]:
        expected[(row["preceptor_name"], row["month"], row["work_type"])] += row["no_of_shifts"]
    try:
        for row in scan["daily_by_work_type"]:
            day = CalendarDate.fromisoformat(row["date"])
            count = row["no_of_shifts"]
            if (type(count) is not int or count <= 0 or not row["preceptor_name"] or not row["work_type"]
                    or row["educational_hours"] != count * TEACHING_HOURS_PER_STUDENT_SHIFT
                    or not isinstance(row["source_sites"], list) or not row["source_sites"]):
                raise ValueError("invalid daily aggregate")
            actual[(row["preceptor_name"], day.replace(day=1).isoformat(), row["work_type"])] += count
    except (KeyError, TypeError, ValueError):
        raise OPDArchiveError("The daily teaching details are incomplete. Refresh the archived OPDs before generating a report.") from None
    if dict(actual) != dict(expected):
        raise OPDArchiveError("Daily teaching counts do not match monthly work-type totals. Refresh the archived OPDs; no partial report was generated.")


def teaching_filter_date_range(scan, period):
    """Project one full archive scan into one inclusive, user-labeled period.

    Filter the *actual assignment date* before computing month/work-type totals.
    July boundaries and twelve-month lengths play no role. Unlinked student-date
    groups are filtered too, without learner names or changing the original scan.
    Archive source/diagnostic
    counts still describe whole files and are explicitly labeled as such in the UI
    and ZIP notes. Always pass the full scan, not an already filtered projection.
    """
    if not isinstance(period, ReportingPeriod):
        raise OPDArchiveError("Choose a valid reporting label, start date and end date.")
    if scan.get("reporting_period") is not None:
        raise OPDArchiveError("Choose the reporting dates from the full archive scan, not an already filtered report.")
    teaching_require_date_range_data(scan)
    selected = [dict(row, source_sites=list(row["source_sites"])) for row in scan["daily_by_work_type"]
                if period.start.isoformat() <= row["date"] <= period.end.isoformat()]
    overall, typed, sites = Counter(), Counter(), defaultdict(set)
    today = teaching_local_today().isoformat()
    for row in selected:
        month = CalendarDate.fromisoformat(row["date"]).replace(day=1).isoformat()
        key = (row["preceptor_name"], month, row["work_type"])
        overall[(row["preceptor_name"], month)] += row["no_of_shifts"]
        typed[key] += row["no_of_shifts"]
        sites[key].update(row["source_sites"])
    # This integer is a compatibility grouping key for the existing report
    # builders. It is NOT a July-based academic-year classification.
    group_id = period.start.year
    monthly = [{"preceptor_name": name, "academic_year": period.label,
                "academic_start_year": group_id, "month": month,
                "no_of_shifts": count, "educational_hours": count * TEACHING_HOURS_PER_STUDENT_SHIFT}
               for (name, month), count in sorted(overall.items(), key=lambda pair: (teaching_name_key(pair[0][0]), pair[0][1]))]
    monthly_types = [{"preceptor_name": name, "academic_year": period.label,
                      "academic_start_year": group_id, "month": month, "work_type": work_type,
                      "no_of_shifts": count, "educational_hours": count * TEACHING_HOURS_PER_STUDENT_SHIFT,
                      "source_sites": sorted(sites[(name, month, work_type)])}
                     for (name, month, work_type), count in sorted(
                         typed.items(), key=lambda pair: (teaching_name_key(pair[0][0]), pair[0][1], teaching_work_type_sort(pair[0][2])))]
    names = {row["preceptor_name"] for row in monthly}
    selected_sites = {site for row in selected for site in row["source_sites"]}
    result = dict(scan)
    result.update({
        "reporting_period": period.as_dict(),
        "daily_by_work_type": selected,
        "monthly": monthly,
        "monthly_by_work_type": monthly_types,
        "work_type_conflicts": [dict(row, academic_year=period.label)
                                for row in scan.get("work_type_conflicts", [])
                                if period.start.isoformat() <= row["date"] <= period.end.isoformat()],
        "unresolved_preceptor_labels": [name for name in scan.get("unresolved_preceptor_labels", []) if name in names],
        "name_variants": {name: list(values) for name, values in scan.get("name_variants", {}).items() if name in names},
        "site_work_type_mapping": {site: label for site, label in scan.get("site_work_type_mapping", {}).items()
                                   if site in selected_sites},
        "future_assignments_in_period": sum(row["no_of_shifts"] for row in selected if row["date"] > today),
    })
    if "learner_reach_version" in scan:
        filter_reach_dates(scan, result, period)
    if "student_continuity_version" in scan or "student_day_groups" in scan:
        filter_student_continuity_dates(scan, result, period)
    teaching_require_date_range_data(result)
    return result
